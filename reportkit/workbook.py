"""Test case workbook from the project's template (KBKT layout by default): one sheet, the test cases.

Claude writes a short builder script in the workspace:

    from reportkit import profile as P, workbook as W
    prof = P.load(); ws = prof.workspace("1E_117")
    wb = W.Workbook(prof, ws, name="Trang tổng hợp CBTT nội bộ", screen="[# 1E_117] ... - IDS > Trang chủ",
                    ticket="AI-23", ascii_name="TrangTongHopCongBoThongTinNoiBo", version=1)
    wb.chapter("Chức năng 1: ..."); wb.pre("1. Đăng nhập ...")
    wb.cat("Giao diện"); wb.sub("Giao diện chung")
    wb.tc("Kiểm tra ...", ["Chọn Năm = 2026", "Bấm [Tìm kiếm]"], ["Ý mong muốn 1", "Ý mong muốn 2"],
          basis="Căn cứ: ...", status="F", actual=["Sai.", "Ví dụ: ...", "Nguyên nhân: ..."])
    wb.finish(run_date="28/09/2026", run_note="Thời gian: ...\\nNgười thực hiện: ...\\nBản build: ...")
    wb.save()

A list is written one item per line: steps are numbered "1. ", expected lines get "- ", and in 'actual' the first
item (the verdict) stays as is and the others get "- ".

Coverage (enforced by save()): every requirement must have a case or a stated reason, otherwise no file is written.
Requirements are read from the workspace:
  - analysis.md: every rule line that starts with its id ("R1. ...", "- R2: ...", "| R3 | ..."), id "R1";
  - inputs/transactions.json (written by `rt start` from the function list's transaction sheet): id "T1", "T2"...;
  - probe.md: every "[FOUND] <trap>:" line, id "trap:<trap>";
  - the checklist areas of test-areas.md (COVERAGE_AREAS, or profile workbook.coverage.areas), id "area:<key>".
    wb.tc(..., covers=["R4", "trap:lang_twins", "area:bo_loc"])
    wb.waive("R9", "Không viết case: chức năng chưa nối vào màn hình (đưa vào hand-off)")
save() writes <workspace>/coverage.md (requirement -> cases or reason) for the hand-off.
`rt build <code> --check` runs the builder into <workspace>/_check.xlsx with every gate, without using up a version.

Profile block (defaults shown are the KBKT template):
  workbook:
    template: "{tools_dir}/Template Testcase/KBKT_Template.xlsx"
    out_dir: "{tools_dir}"
    file_name: "{ticket}_{code_nodash}_{ascii_name}_v{version}.xlsx"
    first_row: 12
    style_rows: {chapter: 12, pre: 13, cat: 14, sub: 15, grp: 24, case: 16}
    header_cells: {screen: D2, code: D3}
    run_header_row: 11            # "Lần N" header row; E=run 1, F=run 2, G=run 3
    font: "Times New Roman"
"""
import copy
import json
import os
import re

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

GREEN = PatternFill("solid", fgColor="FFC6EFCE"); RED = PatternFill("solid", fgColor="FFFFC7CE")
YELLOW = PatternFill("solid", fgColor="FFFFEB9C")
STATUS_FILL = {"P": GREEN, "F": RED, "PE": YELLOW}
VERDICT = {"P": ("Đạt.",), "F": ("Sai.",), "PE": ("Cần BA xác nhận.",),
           None: ("Chưa thực hiện", "Không áp dụng")}          # how column H must start, by column E

# Readability gate (`rt build`): what a Tester / BA reader should not meet in B, C, D, H. The security block is skipped
# (its steps carry payloads), and column J is only checked for run ids and length: its 'Kỹ thuật:' line is for DEV.
LINE_MAX = 160                                     # characters per line of B, D, H, J: longer means several ideas
_SYMBOLS = re.compile(r"[Σ∑→⇒≥≤×|]")
_JARGON = re.compile(r"\b(?:vd|v\.v|API|view|dedup|NULL|NaN|MATCH|NM|DIFF|CTE|request|query|ô so sánh)\b|"
                     r"\.(?:ts|html|java|sql|py|xml):\d")
_NAMES = re.compile(r"\b[A-Z][A-Z0-9]*_[A-Z0-9_]+\b")      # TABLE_NAME / COLUMN_NAME (not checked in C: steps may type a code)
_RUN_ID = re.compile(r"\b(?:[DUVP]\d{2}|run \d{8}-\d{6})\b")
_BUG_ID = re.compile(r"BUG-\d+")

# Coverage gate: the checklist of references/test-areas.md, one key per item a report screen is expected to cover.
# A key that does not fit the screen is waived with its reason (wb.waive), never silently dropped.
COVERAGE_AREAS = {
    "giao_dien": "Bố cục màn hình: tiêu đề, trường lọc, nút, cột lưới, không lỗi khi mở",
    "dropdown": "Từng dropdown: giá trị, đã xoá, tìm nhanh trên 10 giá trị, thứ tự (Quy chuẩn II.5.1)",
    "tieng_anh_nhan": "English: tiêu đề, nhãn, nút, giá trị dropdown (Quy chuẩn I.1)",
    "tieng_anh_du_lieu": "English: dữ liệu lưới và file giữ tiếng Việt (Quy chuẩn I.1)",
    "file_excel": "File Excel: tiêu đề, cột, định dạng giá trị (Quy chuẩn III.1)",
    "bo_loc": "Từng bộ lọc trên màn hình lọc đúng và đủ",
    "khong_co_du_lieu": "Không có dữ liệu: lưới báo Không có dữ liệu, file vẫn có tiêu đề (Quy chuẩn III.2, III.3)",
    "luoi_va_file": "Lưới và file cùng dòng, cùng giá trị cho cùng bộ lọc",
    "pham_vi_cong_ty": "Phạm vi công ty: đã xoá, IPO, thu hồi; dropdown, lưới, file cùng phạm vi",
    "phan_quyen": "Phân quyền dữ liệu: dropdown, lưới, file (Quy chuẩn I.3)",
    "doi_chieu_man_hinh_khac": "Đối chiếu với màn hình hoặc báo cáo khác dùng cùng dữ liệu",
    "xss": "An toàn thông tin: XSS (bộ lọc, đường dẫn, dữ liệu trả về)",
    "sql_injection": "An toàn thông tin: SQL Injection",
}
_RULE_LINE = re.compile(r"^\s*(?:[-*]\s+|\|\s*)?\**(R\d+[a-z]?)\**\s*[.:)|]\s*(.*)$")
_TRAP_LINE = re.compile(r"^\s*-\s*\[FOUND\]\s*([A-Za-z0-9_]+)\s*:\s*(.*)$")


def requirements(ws_dir, areas=None):
    """{id: (source, text)} of everything the workbook must cover: rule ids of analysis.md, the function's transactions,
    FOUND traps of probe.md, checklist areas. analysis.md must exist and number its rules (R1, R2, ...)."""
    path = os.path.join(ws_dir, "analysis.md")
    if not os.path.exists(path):
        raise SystemExit("coverage: %s is missing. Write it first (SKILL.md step 3): every rule on its own line, "
                         "numbered R1, R2, ..." % path)
    req = {}
    with open(path, encoding="utf-8") as f:
        for line in f:
            m = _RULE_LINE.match(line)
            if m and m.group(1) not in req:
                req[m.group(1)] = ("analysis.md", m.group(2).strip()[:140])
    if not req:
        raise SystemExit("coverage: analysis.md has no numbered rule. Start each rule line with its id: 'R1. ...'")
    tx = os.path.join(ws_dir, "inputs", "transactions.json")
    if os.path.exists(tx):
        with open(tx, encoding="utf-8") as f:
            for t in json.load(f).get("items") or []:
                req[t["id"]] = ("transaction", t["text"][:140])
    probe = os.path.join(ws_dir, "probe.md")
    if os.path.exists(probe):
        with open(probe, encoding="utf-8") as f:
            for line in f:
                m = _TRAP_LINE.match(line)
                if m and "trap:" + m.group(1) not in req:
                    req["trap:" + m.group(1)] = ("probe.md", m.group(2).strip()[:140])
    for k, v in (COVERAGE_AREAS if areas is None else areas).items():
        req["area:" + k] = ("test-areas.md", v)
    return req


def _as_list(v):
    return [] if v is None else [v] if isinstance(v, str) else list(v)


def set_text(cell, value):
    """Text starting with '=' stays text (openpyxl would store a formula and Excel would 'repair' the file)."""
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def _text_width(s):
    """Width of s in column-width units for a 12pt Times New Roman font (capitals and digits are wider)."""
    return sum(1.35 if ch.isupper() else 1.1 if (ch.isdigit() or ch in "_%@#&") else 0.55 if ch in " .,:;'|!il" else 1.0 for ch in s)


def _wrapped_lines(text, avail):
    """Lines Excel needs for text in a cell `avail` units wide: word wrap (Excel also breaks after '-'),
    pieces longer than a line are broken."""
    import re
    n = 0
    for para in str(text).split("\n"):
        line, lines = 0.0, 1
        for word in para.split(" "):
            for k, piece in enumerate(re.findall(r"[^-]*-|[^-]+", word) or [""]):
                gap = 0.55 if (line and k == 0) else 0.0      # a space only before the first piece of a word
                w = _text_width(piece)
                if w > avail:                              # a long identifier breaks across lines
                    if line:
                        lines += 1
                    full, rest = divmod(w, avail)
                    lines += int(full) - (0 if rest else 1)
                    line = rest
                elif line and line + gap + w > avail:
                    lines += 1; line = w
                else:
                    line += gap + w
        n += lines
    return n


def fit_rows(sheet, rows, cols=None, line_pt=15.75, pad=4.0, min_pt=15.75, max_pt=409.0):
    """Set each row's height so wrapped text shows in full. openpyxl cannot autofit, and the template's
    fixed 15.75pt heights would clip the text to one line. Height = the most wrapped lines of any cell
    in the row (word wrap simulated from the column width), checked against Excel's own AutoFit."""
    widths = {}
    for r in rows:
        most = 1
        for c in (cols or range(1, sheet.max_column + 1)):
            v = sheet.cell(r, c).value
            if v is None or v == "":
                continue
            L = openpyxl.utils.get_column_letter(c)
            if L not in widths:
                widths[L] = sheet.column_dimensions[L].width or 8.43
            most = max(most, _wrapped_lines(v, max(1.0, widths[L] - 1.2)))
        sheet.row_dimensions[r].height = max(min_pt, min(max_pt, most * line_pt + pad))


class Workbook(object):
    def __init__(self, prof, ws, name, screen, ticket="", ascii_name="", version=1, out=None, sheet_title=None):
        cfg = prof.get("workbook", {}) or {}
        self.cfg = cfg
        self.space = ws
        tpl = prof.expand(cfg.get("template", "{tools_dir}/Template Testcase/KBKT_Template.xlsx"))
        if not os.path.exists(tpl):
            raise SystemExit("Workbook template not found: %s (profile workbook.template)" % tpl)
        fname = cfg.get("file_name", "{ticket}_{code_nodash}_{ascii_name}_v{version}.xlsx").format(
            ticket=ticket, code=ws.code, code_nodash=ws.code.replace("_", ""), ascii_name=ascii_name or ws.code, version=version)
        check_out = os.environ.get("RK_WB_OUT")                  # `rt build --check`: a scratch file in the workspace
        self.out = out or check_out or os.path.join(prof.expand(cfg.get("out_dir", "{tools_dir}")), fname)
        if check_out and self.out == check_out and os.path.exists(check_out):
            os.remove(check_out)
        if os.path.exists(self.out):
            raise SystemExit("%s already exists. Never overwrite a delivered version - use version=%d." % (self.out, version + 1))
        f = cfg.get("font", "Times New Roman")
        self.FONT = Font(name=f, size=12); self.FONT_B = Font(name=f, size=12, bold=True)
        self.wb = openpyxl.load_workbook(tpl)
        sh = self.ws = self.wb.worksheets[0]
        sh.title = (sheet_title or ("%s %s" % (ws.code, name)))[:31]
        rows = cfg.get("style_rows", {"chapter": 12, "pre": 13, "cat": 14, "sub": 15, "grp": 24, "case": 16})
        cap = lambda r: [copy.copy(sh.cell(r, c)._style) for c in range(1, 11)]
        self.st = {k: cap(r) for k, r in rows.items()}
        self.first = int(cfg.get("first_row", 12))
        self.cases = []                       # (row, purpose) of every tc()
        self.covered = {}                     # requirement id -> [row, ...]
        self.waived = {}                      # requirement id -> reason
        cov = cfg.get("coverage") or {}
        self.areas = cov.get("areas") if isinstance(cov.get("areas"), dict) else None
        self.pre_h = sh.row_dimensions[rows.get("pre", 13)].height
        for m in [str(m) for m in sh.merged_cells.ranges if m.min_row >= self.first]:
            sh.unmerge_cells(m)
        sh.delete_rows(self.first, sh.max_row)
        # delete_rows() leaves the template's row properties behind: from row ~101 the KBKT template has
        # hidden, grouped rows, so every case written there was invisible in Excel (1E_119, 1E_126, 2E_54 v2...).
        for r in [r for r in sh.row_dimensions if r >= self.first]:
            del sh.row_dimensions[r]
        sh._images = []                       # the KBKT template carries thousands of stray images
        sh.data_validations.dataValidation = []
        hc = cfg.get("header_cells", {"screen": "D2", "code": "D3"})
        sh[hc["screen"]] = screen; sh[hc["code"]] = ws.code
        self.code_cell = hc["code"]
        self.row = self.first

    def _put(self, kind, b=None, c=None, d=None, j=None, merge=False):
        sh, r = self.ws, self.row
        for col in range(1, 11):
            sh.cell(r, col)._style = copy.copy(self.st[kind][col - 1])
        col_d, row_d = self.code_cell[0], int(self.code_cell[1:])
        hdr = self.first - 1
        sh.cell(r, 1).value = '=IF(AND(D%d="",D%d=""),"",$%s$%d&"_"&ROW()-%d-COUNTBLANK($D$%d:D%d))' % (r, r, col_d, row_d, hdr - 1, hdr, r)
        for col, v in ((2, b), (3, c), (4, d), (10, j)):
            if v is not None:
                set_text(sh.cell(r, col), v)
        if merge:
            sh.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10)
        if kind == "case":
            for col in (2, 3, 4, 8, 9, 10):
                sh.cell(r, col).alignment = Alignment(wrap_text=True, vertical="top")
        self.row += 1
        return r

    def chapter(self, t): self._put("chapter", t, merge=True)
    def cat(self, t): self._put("cat", t, merge=True)
    def sub(self, t): self._put("sub", t, merge=True)
    def grp(self, t): self._put("grp", t, merge=True)

    def pre(self, t):
        r = self._put("pre", t, merge=True); self.ws.row_dimensions[r].height = self.pre_h

    def waive(self, req_id, reason):
        """A requirement with no case, and why (it goes to coverage.md and the hand-off, not to the sheet)."""
        if not reason or len(str(reason).strip()) < 10:
            raise ValueError("waive(%r): give the reason in a sentence" % req_id)
        self.waived[req_id] = str(reason).strip()

    def tc(self, purpose, steps, expected, basis=None, status=None, actual=None, bug=None, run=1, covers=None):
        """steps / expected / actual / basis: a string, or a list written one item per line (see the module doc).
        bug: optional, the tracker's bug id once the Tester has logged it.
        covers: requirement ids this case covers ("R3", "trap:lang_twins", "area:bo_loc"), see the module doc."""
        steps = _lines(steps, number=True)
        expected = _lines(expected, bullet=True)
        if isinstance(actual, (list, tuple)):
            actual = "\n".join([str(actual[0])] + ["- %s" % x for x in actual[1:]])
        basis = _lines(basis)
        if status is None and not actual:
            actual = "Chưa thực hiện"
        if status in ("P", "F", "PE") and actual in (None, "", "P", "F", "PE"):
            raise ValueError("tc(%r): 'actual' must describe what happened, not just the status" % purpose)
        r = self._put("case", purpose, steps, expected, basis)
        set_text(self.ws.cell(r, 8), actual)
        if status:
            c = self.ws.cell(r, 4 + run); c.value = status; c.fill = STATUS_FILL[status]
        if bug:
            self.ws.cell(r, 9).value = bug
        self.cases.append((r, purpose))
        for c in _as_list(covers):
            self.covered.setdefault(c, []).append(r)
        return r

    def coverage(self):
        """(requirements, missing ids, unknown ids used in covers / waive)."""
        req = requirements(self.space.dir, self.areas)
        used = set(self.covered) | set(self.waived)
        missing = [k for k in req if k not in used]
        unknown = sorted(k for k in used if k not in req)
        return req, missing, unknown

    def _write_coverage(self, req, missing, unknown):
        sh, hdr = self.ws, self.first - 1
        code = sh[self.code_cell].value
        purpose = dict(self.cases)
        n_case = sum(1 for k in req if k in self.covered)
        n_wv = sum(1 for k in req if k not in self.covered and k in self.waived)
        lines = ["# Coverage %s: %d requirements, %d by cases, %d without a case (reason), %d missing" %
                 (code, len(req), n_case, n_wv, len(missing)), "",
                 "| Requirement | Source | What | Covered by / reason |", "|---|---|---|---|"]
        for k, (src, text) in req.items():
            if k in self.covered:
                how = "; ".join("%s %s" % (_case_id(sh, hdr, code, r), purpose.get(r, "")) for r in self.covered[k])
            elif k in self.waived:
                how = "Không có case: " + self.waived[k]
            else:
                how = "**MISSING**"
            lines.append("| %s | %s | %s | %s |" % (k, src, text.replace("|", "/"), how.replace("|", "/")))
        if unknown:
            lines += ["", "Unknown ids used in covers / waive (typo?): " + ", ".join(unknown)]
        with open(os.path.join(self.space.dir, "coverage.md"), "w", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")
        return lines[0]

    def finish(self, run_date, run_note, author="", run=1):
        sh, last, f = self.ws, self.row - 1, self.first
        sh["D4"] = '=COUNTIF($E$%d:$E$%d,"P")' % (f, last)
        sh["D5"] = '=COUNTIF($E$%d:$E$%d,"F")' % (f, last)
        sh["D6"] = '=COUNTIF($E$%d:$E$%d,"PE")' % (f, last)
        sh["D7"] = "=D8-D4-D5-D6"
        sh["D8"] = "=COUNTA($D$%d:$D$%d)" % (f, last)
        hr = int(self.cfg.get("run_header_row", f - 1))
        cell = sh.cell(hr, 4 + run); cell.value = "Lần %d\n%s" % (run, run_date)
        cell.comment = Comment(run_note, author or "reportkit")
        sh.cell(hr - 1, 8).comment = Comment(run_note, author or "reportkit")
        sh.column_dimensions["H"].width = 55; sh.column_dimensions["J"].width = 40
        dv = DataValidation(type="list", formula1='"P,F,PE"', allow_blank=True); dv.add("E%d:G%d" % (f, last))
        sh.add_data_validation(dv)
        fit_rows(sh, [r for r, _ in self.cases], cols=(2, 3, 4, 8, 9, 10))
        return last

    def save(self):
        req, missing, unknown = self.coverage()
        head = self._write_coverage(req, missing, unknown)
        if missing or unknown:
            msg = [head.lstrip("# "), "No file written. Add a case (tc(..., covers=[...])) or wb.waive(id, reason) for:"]
            msg += ["  %s (%s): %s" % (k, req[k][0], req[k][1]) for k in missing]
            if unknown:
                msg.append("  unknown ids in covers / waive (typo?): " + ", ".join(unknown))
            msg.append("Details: %s" % os.path.join(self.space.dir, "coverage.md"))
            raise SystemExit("\n".join(msg))
        self.wb.active = 0
        os.makedirs(os.path.dirname(self.out), exist_ok=True)
        self.wb.save(self.out)
        return self.out


def _lines(v, number=False, bullet=False):
    """A list -> one item per line ('1. ' numbered or '- ' bulleted); a single item or a string is kept as is."""
    if not isinstance(v, (list, tuple)):
        return v
    if len(v) == 1:
        return str(v[0])
    if number:
        return "\n".join("%d. %s" % (i, x) for i, x in enumerate(v, 1))
    return "\n".join(("- %s" % x) if bullet else str(x) for x in v)


def _case_id(sh, hdr, code, r):
    """Value of the column-A formula of row r: code_N, N = non-blank 'Kết quả mong muốn' cells from the header row to r."""
    return "%s_%d" % (code, sum(1 for rr in range(hdr, r + 1) if sh.cell(rr, 4).value not in (None, "")))


def readability(path, first_row=12, code_cell="D3"):
    """Quality gate on the wording of the test case sheet, for a Tester / BA reader.
    Returns [(case id, column letter, problem)]: H not starting with the verdict of column E, symbols or internal names
    in B / C / D / H, a line longer than LINE_MAX, several bug ids in one case, run ids in J."""
    sh = openpyxl.load_workbook(path).worksheets[0]
    code, hdr, sec, out = sh[code_cell].value, first_row - 1, None, []
    for r in range(first_row, sh.max_row + 1):
        if sh.cell(r, 4).value is None:
            if sh.cell(r, 2).value in ("Giao diện", "Chức năng", "An toàn thông tin"):
                sec = sh.cell(r, 2).value
            continue
        cid = _case_id(sh, hdr, code, r)
        txt = {L: str(sh.cell(r, c).value or "") for L, c in (("B", 2), ("C", 3), ("D", 4), ("H", 8), ("I", 9), ("J", 10))}
        want = VERDICT.get(sh.cell(r, 5).value, VERDICT[None])
        if not txt["H"].startswith(want):
            out.append((cid, "H", "must start with '%s' (column E = %s)" % ("' / '".join(want), sh.cell(r, 5).value or "empty")))
        if len(_BUG_ID.findall(txt["I"])) > 1:
            out.append((cid, "I", "several bugs in one case: split it, one case per bug"))
        if _RUN_ID.search(txt["J"]):
            out.append((cid, "J", "run id (%s): means nothing to the reader, drop it" % _RUN_ID.search(txt["J"]).group(0)))
        for L in ("B", "D", "H", "J"):
            if any(len(x) > LINE_MAX for x in txt[L].split("\n")):
                out.append((cid, L, "a line over %d characters: one idea per line" % LINE_MAX))
        if sec == "An toàn thông tin":
            continue
        for L in ("B", "C", "D", "H"):
            m = _SYMBOLS.search(txt[L]) or _JARGON.search(txt[L]) or (L != "C" and _NAMES.search(txt[L]))
            if m:
                out.append((cid, L, "'%s': write it in words, as the screen names it (technical names go to 'Kỹ thuật:' in J)" % m.group(0)))
    return out


_CONV_REF = re.compile(r"Quy ước kiểm thử nhóm\s+([^\s,;]+)\s+mục\s+([^\s,;.)]+)")


def convention_refs(path, first_row=12, code_cell="D3"):
    """([(case id, group, item)] cited in column J, whether a non-case row (chapter / precondition) says who set the
    group conventions). Cited as 'Căn cứ: Quy ước kiểm thử nhóm 1D.I mục Q2'."""
    sh = openpyxl.load_workbook(path).worksheets[0]
    code, hdr, refs, pre_ok = sh[code_cell].value, first_row - 1, [], False
    for r in range(first_row, sh.max_row + 1):
        if sh.cell(r, 4).value is None:
            pre_ok = pre_ok or "Quy ước kiểm thử nhóm" in str(sh.cell(r, 2).value or "")
            continue
        for m in _CONV_REF.finditer(str(sh.cell(r, 10).value or "")):
            refs.append((_case_id(sh, hdr, code, r), m.group(1), m.group(2)))
    return refs, pre_ok


def hidden_rows(path, first_row=12):
    """Rows of the test case sheet that Excel will not show (hidden rows or collapsed groups) - quality gate."""
    sh = openpyxl.load_workbook(path).worksheets[0]
    return [r for r in range(first_row, sh.max_row + 1)
            if sh.row_dimensions[r].hidden and (sh.cell(r, 2).value or sh.cell(r, 4).value)]


def tally(path, first_row=12):
    """Counts per section + rows whose 'Kết quả hiện tại' is empty or only a status (quality gate)."""
    import collections
    sh = openpyxl.load_workbook(path).worksheets[0]
    sec, per, bad = None, collections.defaultdict(collections.Counter), []
    for r in range(first_row, sh.max_row + 1):
        b = sh.cell(r, 2).value
        if sh.cell(r, 4).value is None and b in ("Giao diện", "Chức năng", "An toàn thông tin"):
            sec = b
        if sh.cell(r, 4).value:
            st = sh.cell(r, 5).value or "Chưa thực hiện"
            per[sec][st] += 1
            if not sh.cell(r, 8).value or sh.cell(r, 8).value in ("P", "F", "PE"):
                bad.append(r)
    return {k: dict(v) for k, v in per.items()}, bad
