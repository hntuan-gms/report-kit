"""Read an exported Excel file as the user sees it (number / date formats applied)."""
import datetime as _dt
import re


def shown(cell):
    """Text a user sees in Excel for this cell. Never str() a cell value: 2.0 with '#,##0' shows '2'."""
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, _dt.datetime):
        f = (cell.number_format or "").lower()
        if f.startswith("dd/mm/yy"):
            return v.strftime("%d/%m/%Y")
        if f.startswith("mm/dd"):
            return v.strftime("%m/%d/%Y")
        return v.strftime("%d/%m/%Y") + " [fmt %s]" % cell.number_format
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        fmt = cell.number_format or "General"
        m = re.match(r"^#,##0(?:\.(0*)(#*))?(?:;.*)?$", fmt)
        if m:
            mn = len(m.group(1) or ""); mx = mn + len(m.group(2) or "")
            t = ("{:,.%df}" % mx).format(v)
            while "." in t and len(t.split(".")[1]) > mn and t.endswith("0"):
                t = t[:-1]
            return t.rstrip(".") if mn == 0 and t.endswith(".") else t
        m = re.match(r"^0(?:\.(0+))?$", fmt)
        if m:
            return ("{:.%df}" % len(m.group(1) or "")).format(v)
        if fmt.endswith("%"):
            d = len(fmt.rstrip("%").split(".")[1]) if "." in fmt else 0
            return ("{:.%df}" % d).format(v * 100) + "%"
        return "%.10g" % v
    return str(v)


def read(path, header_marker="STT", marker_col_max=6, sheet=None):
    """{"meta": {label: value above the table}, "header_row", "cells": {(r, c): shown text}, "types", "max_row", "max_col", "ws"}."""
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet] if sheet else wb.active
    hr = None
    for r in range(1, 40):
        for c in range(1, marker_col_max + 1):
            if str(ws.cell(r, c).value or "").strip().upper() == header_marker.upper():
                hr = r; break
        if hr:
            break
    meta = {}
    for r in range(1, (hr or 10)):
        vals = [(c, ws.cell(r, c).value) for c in range(1, 10) if ws.cell(r, c).value not in (None, "")]
        if len(vals) >= 2:
            meta[str(vals[0][1]).strip().rstrip(":")] = vals[1][1]
    merged = {}
    for m in ws.merged_cells.ranges:
        v = ws.cell(m.min_row, m.min_col).value
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                merged[(r, c)] = v
    cells, types = {}, {}
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            x = ws.cell(r, c)
            if (r, c) in merged and x.value is None:
                cells[(r, c)] = "" if merged[(r, c)] is None else str(merged[(r, c)])
            else:
                cells[(r, c)] = shown(x)
            types[(r, c)] = (x.data_type, x.number_format)
    return {"meta": meta, "header_row": hr, "cells": cells, "types": types, "max_row": ws.max_row,
            "max_col": ws.max_column, "ws": ws, "path": path}


def headers(sheet, n_rows=1, sep=" > "):
    hr = sheet["header_row"] or 1; out = {}
    for c in range(1, sheet["max_col"] + 1):
        parts = []
        for r in range(hr, hr + n_rows):
            t = sheet["cells"].get((r, c), "").strip()
            if t and (not parts or parts[-1] != t):
                parts.append(t)
        if parts:
            out[c] = sep.join(parts)
    return out


def records(sheet, n_header_rows=1, key_col=1):
    """List reports: one dict per data row below the header rows (stops at the first fully blank row)."""
    hdr = headers(sheet, n_header_rows); rows = []
    for r in range((sheet["header_row"] or 1) + n_header_rows, sheet["max_row"] + 1):
        if not sheet["cells"].get((r, key_col), "").strip():
            if all(not sheet["cells"].get((r, c), "").strip() for c in hdr):
                break
            continue
        rows.append({h: sheet["cells"].get((r, c), "") for c, h in hdr.items()})
    return rows


def grid_by_key(sheet, key_col=1, first_value_col=None, n_header_rows=1):
    """Fixed-row reports: {(row key, column header): shown text}."""
    hdr = headers(sheet, n_header_rows); out = {}
    for r in range((sheet["header_row"] or 1) + n_header_rows, sheet["max_row"] + 1):
        k = sheet["cells"].get((r, key_col), "").strip()
        if not k:
            continue
        try:
            k = str(int(float(k)))
        except ValueError:
            pass
        for c, h in hdr.items():
            if first_value_col and c < first_value_col:
                continue
            out[(k, h)] = sheet["cells"].get((r, c), "")
    return out


def missing_borders(sheet, rows, cols):
    ws = sheet["ws"]; bad = []
    for r in rows:
        for c in cols:
            b = ws.cell(r, c).border
            if not all(s.style for s in (b.left, b.right, b.top, b.bottom)):
                bad.append(ws.cell(r, c).coordinate)
    return bad


def std_number(x, decimals=None):
    """Display a number as 1,234.56 (thousands ',', decimal '.')."""
    x = float(x)
    if decimals is not None:
        return ("{:,.%df}" % decimals).format(x)
    if x == int(x):
        return "{:,}".format(int(x))
    return "{:,.6f}".format(x).rstrip("0").rstrip(".")
