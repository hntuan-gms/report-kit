"""Find and dump the test inputs (function list row, standard, previous workbook) into text.

Everything is driven by the profile's `inputs:` block, e.g.

inputs:
  function_list:
    glob: "*Danh sách chức năng*.xlsx"      # under tools_dir
    sheet: 2                                # 1-based index or sheet name (default: the first sheet)
    header_contains: "Mã Jira"              # a cell text that identifies the header row
    code_columns: ["New UC Name"]           # where the report code is looked up (default: every column)
    columns: {jira: "Mã Jira", name: "New UC Name", system: "Phân hệ", menu: "Menu/tab (cấp thấp nhất)", ...}
    ticket: {column: testcase, regex: "^([A-Z]+-\\d+)_"}   # ticket taken from another column when there is no ticket column
  transactions:                             # optional: what each function must offer, one row per user action
    sheet: "Chi tiết transaction"           # in the function list file (or `glob` for another file)
    system: "IDS"                           # the sheet lists only this Phân hệ (its Stt restarts for the others)
    columns: {stt: "Stt", text: "Mô tả yêu cầu", kind: "Phân loại"}
  standard: {glob: "**/SRS_Quy_chuan_chung*.docx", summary: rules/standard.md}
  previous_workbooks: "*{code_nodash}*.xlsx"

Column headers match exactly (case-insensitive) first, then as a substring.
"""
import glob
import os
import re
import zipfile


# ------------------------------------------------------------------ dumpers
def dump_xlsx(path, out=None, max_rows=3000, with_comments=True, images_dir=None):
    import openpyxl
    lines = []
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        lines.append("===== SHEET: %s dims=%s" % (ws.title, ws.dimensions))
        n = 0
        for row in ws.iter_rows():
            vals = [(c.coordinate, c.value) for c in row if c.value is not None and str(c.value).strip() != ""]
            if vals:
                lines.append(" | ".join("%s=%s" % (k, str(v).replace("\n", " / ")) for k, v in vals))
                n += 1
                if n >= max_rows:
                    lines.append("...TRUNCATED"); break
    if with_comments:
        wb2 = openpyxl.load_workbook(path)
        for ws in wb2.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.comment:
                        lines.append("##### COMMENT %s!%s on [%s]: %s" % (ws.title, c.coordinate, str(c.value)[:200],
                                                                        c.comment.text.replace("\n", " / ")))
    imgs = extract_images(path, images_dir or (os.path.splitext(out)[0] + "_images" if out else None))
    return _emit(lines, out), imgs


def dump_docx(path, out=None, images_dir=None):
    import docx
    from docx.oxml.ns import qn
    lines = []
    for child in docx.Document(path).element.body.iterchildren():
        if child.tag == qn("w:p"):
            t = "".join(x.text or "" for x in child.iter(qn("w:t")))
            if t.strip():
                lines.append(t)
        elif child.tag == qn("w:tbl"):
            for tr in child.iter(qn("w:tr")):
                lines.append(" || ".join(" ".join((x.text or "") for x in tc.iter(qn("w:t"))).strip() for tc in tr.iter(qn("w:tc"))))
            lines.append("--- end table")
    imgs = extract_images(path, images_dir or (os.path.splitext(out)[0] + "_images" if out else None))
    return _emit(lines, out), imgs


def dump_any(path, out=None, images_dir=None):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return dump_xlsx(path, out, images_dir=images_dir)
    if ext == ".docx":
        return dump_docx(path, out, images_dir=images_dir)
    with open(path, encoding="utf-8", errors="replace") as f:
        return _emit(f.read().splitlines(), out), []


def extract_images(path, out_dir):
    if not out_dir:
        return []
    saved = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if re.match(r"^(xl|word)/media/", n):
                os.makedirs(out_dir, exist_ok=True)
                p = os.path.join(out_dir, os.path.basename(n))
                with open(p, "wb") as f:
                    f.write(z.read(n))
                saved.append(p)
    return saved


def _emit(lines, out):
    text = "\n".join(lines)
    if out:
        os.makedirs(os.path.dirname(out), exist_ok=True)
        with open(out, "w", encoding="utf-8") as f:
            f.write(text)
    return text


# ------------------------------------------------------------------ lookups
def _tools(profile, pattern):
    if not profile.tools_dir:
        return []
    return [p for p in glob.glob(os.path.join(profile.tools_dir, pattern), recursive=True) if not os.path.basename(p).startswith("~$")]


def _function_rows(profile):
    """(source file, [raw row dict]) of the function list, read from the sheet the profile names."""
    import openpyxl
    fl = profile.get("inputs.function_list", {}) or {}
    hits = _tools(profile, fl.get("glob", "*.xlsx"))
    if not hits:
        raise SystemExit("Function list '%s' not found in %s" % (fl.get("glob"), profile.tools_dir))
    key = (hits[0], os.path.getmtime(hits[0]), repr(fl))
    if key in _ROWS_CACHE:
        return _ROWS_CACHE[key]
    wb = openpyxl.load_workbook(hits[0], data_only=True, read_only=True)     # read-only: only the named sheet is parsed
    sheet = fl.get("sheet")
    if sheet is None:
        ws = wb.worksheets[0]
    elif isinstance(sheet, int):
        if not 1 <= sheet <= len(wb.worksheets):
            raise SystemExit("Function list %s has %d sheets, profile asks for sheet %d" % (hits[0], len(wb.worksheets), sheet))
        ws = wb.worksheets[sheet - 1]
    elif sheet in wb.sheetnames:
        ws = wb[sheet]
    else:
        raise SystemExit("Sheet '%s' not found in %s (sheets: %s)" % (sheet, hits[0], ", ".join(wb.sheetnames)))
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    title = ws.title
    wb.close()
    marker = fl.get("header_contains", "")
    hdr_i = next((i for i, r in enumerate(grid[:30]) if any(marker in str(v or "") for v in r[:30])), None)
    if hdr_i is None:
        raise SystemExit("Header row containing '%s' not found in %s, sheet '%s'" % (marker, hits[0], title))
    hdr = {c: " ".join(str(v).split()) for c, v in enumerate(grid[hdr_i]) if v not in (None, "")}
    rows = [{hdr[c]: (r[c] if c < len(r) else None) for c in hdr} for r in grid[hdr_i + 1:]]
    _ROWS_CACHE[key] = ("%s [%s]" % (hits[0], title), rows)
    return _ROWS_CACHE[key]


_ROWS_CACHE = {}


def _header(raw, head):
    """The raw header that a profile column name designates: exact match first, then substring."""
    head = " ".join(str(head).split()).lower()
    return (next((h for h in raw if h.lower() == head), None)
            or next((h for h in raw if head in h.lower()), None))


def _logical(profile, raw):
    fl = profile.get("inputs.function_list", {}) or {}
    row = {}
    for logical, head in (fl.get("columns") or {}).items():
        col = _header(raw, head)
        row[logical] = raw.get(col) if col else None
    t = fl.get("ticket")
    if isinstance(t, dict) and not row.get("ticket"):
        m = re.search(t.get("regex", r"([A-Z]+-\d+)"), str(row.get(t.get("column")) or ""))
        row["ticket"] = m.group(1) if m else None
    return row


def _has_code(raw, code, columns=None):
    variants = {code, code.replace("_", "")}
    vals = [raw.get(_header(raw, c)) for c in columns] if columns else list(raw.values())
    return any(any(v and re.search(r"(?<![\w])%s(?![\d])" % re.escape(x), str(v)) for x in variants) for v in vals)


def _code_columns(profile):
    return (profile.get("inputs.function_list", {}) or {}).get("code_columns")


def find_function(profile, code):
    """Row of the function list for `code` as {logical column: value}, plus 'system_key' resolved to a profile key."""
    src, rows = _function_rows(profile)
    for raw in rows:
        if _has_code(raw, code, _code_columns(profile)):
            row = {"_raw": {k: v for k, v in raw.items() if v is not None}, "code": code, "source": src}
            row.update(_logical(profile, raw))
            row["system_key"] = profile.system_for_label(row.get("system"))
            return row
    raise SystemExit("Code %s not found in %s" % (code, src))


# ------------------------------------------------------------------ transactions
_ROLES = re.compile(r"^((?:[A-ZĐ]{2,}\s*[;,]\s*)*[A-ZĐ]{2,})")          # "LĐGSĐC; CVGSĐC" / "LĐTT,CVTT"
_SYS_REPLY = re.compile(r"\.\s*(?=Hệ thống\b)")
_TX_CACHE = {}


def _norm(s):
    return " ".join(str(s or "").split()).lower()


def _stt(v):
    s = str(v if v is not None else "").strip()
    return int(float(s)) if re.fullmatch(r"\d+(?:\.0+)?", s) else None


def _transaction_functions(profile):
    """(source, {stt: {"name", "flag", "items": [(text, kind)]}}) of the transaction sheet.
    A row with a number in Stt opens a function (its 'Phân loại' cell holds a change flag: Sửa / Bổ sung), a row with
    another Stt (A, B.I...) is a section heading, a row with an empty Stt is one transaction of the open function."""
    import openpyxl
    tx = profile.get("inputs.transactions") or {}
    fl = profile.get("inputs.function_list", {}) or {}
    hits = _tools(profile, tx.get("glob") or fl.get("glob", "*.xlsx"))
    if not hits:
        raise SystemExit("Transaction sheet: file '%s' not found in %s" % (tx.get("glob") or fl.get("glob"), profile.tools_dir))
    sheet = tx.get("sheet", "Chi tiết transaction")
    key = (hits[0], os.path.getmtime(hits[0]), repr(tx))
    if key in _TX_CACHE:
        return _TX_CACHE[key]
    wb = openpyxl.load_workbook(hits[0], data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise SystemExit("Transaction sheet '%s' not found in %s (sheets: %s)" % (sheet, hits[0], ", ".join(wb.sheetnames)))
    grid = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    wb.close()
    cols = dict({"stt": "Stt", "text": "Mô tả yêu cầu", "kind": "Phân loại"}, **(tx.get("columns") or {}))
    hdr_i, idx = None, {}
    for i, r in enumerate(grid[:30]):
        heads = {_norm(v): c for c, v in enumerate(r) if v not in (None, "")}
        found = {k: next((c for h, c in heads.items() if h == _norm(v)), None)
                 if k == "stt" else next((c for h, c in heads.items() if _norm(v) in h), None) for k, v in cols.items()}
        if found["stt"] is not None and found["text"] is not None:
            hdr_i, idx = i, found
            break
    if hdr_i is None:
        raise SystemExit("Transaction sheet '%s' of %s: no header row with '%s' and '%s'" % (sheet, hits[0], cols["stt"], cols["text"]))
    cell = lambda r, k: r[idx[k]] if idx.get(k) is not None and idx[k] < len(r) else None
    funcs, cur = {}, None
    for r in grid[hdr_i + 1:]:
        stt, text = cell(r, "stt"), cell(r, "text")
        if _stt(stt) is not None:
            flag = cell(r, "kind")
            cur = funcs.setdefault(_stt(stt), {"name": " ".join(str(text or "").split()),
                                               "flag": str(flag).strip() if flag else None, "items": []})
        elif stt not in (None, ""):
            cur = None
        elif cur is not None and text not in (None, "") and str(text).strip():
            cur["items"].append((" ".join(str(text).split()), str(cell(r, "kind") or "").strip() or None))
    _TX_CACHE[key] = ("%s [%s]" % (hits[0], sheet), funcs)
    return _TX_CACHE[key]


def split_transaction(text):
    """'LĐGSĐC; CVGSĐC xem X. Hệ thống hiển thị Y' -> (['LĐGSĐC', 'CVGSĐC'], 'xem X', 'Hệ thống hiển thị Y')."""
    m = _ROLES.match(text)
    roles = [x for x in re.split(r"\s*[;,]\s*", m.group(1)) if x] if m else []
    rest = text[m.end():].strip() if m else text
    parts = _SYS_REPLY.split(rest, 1)
    return roles, parts[0].strip().rstrip("."), (parts[1].strip() if len(parts) > 1 else "")


def find_transactions(profile, fn):
    """The function's transactions: what each role does and what the system answers (profile inputs.transactions).
    None when the profile has no transaction sheet. Otherwise {"source", "name", "flag", "note", "items": [{"id": "T1",
    "text", "roles", "action", "reply", "kind"}]}; items is empty, with the reason in "note", when the sheet does not
    list the function. Joined on Stt within the sheet's system, and checked on the name ('Mô tả yêu cầu')."""
    tx = profile.get("inputs.transactions")
    if not tx:
        return None
    out = {"source": None, "name": None, "flag": None, "note": None, "items": []}
    if tx.get("system") and _norm(fn.get("system")) != _norm(tx["system"]):
        out["note"] = "the sheet lists only %s functions, this one is %s" % (tx["system"], fn.get("system") or "of no system")
        return out
    src, funcs = _transaction_functions(profile)
    out["source"] = src
    want = _norm(fn.get("description"))
    stt = _stt(fn.get("stt"))
    f = funcs.get(stt) if stt is not None else None
    if f is not None and want and _norm(f["name"]) != want:
        f = None
    if f is None and want:
        same = [g for g in funcs.values() if _norm(g["name"]) == want]
        f = same[0] if len(same) == 1 else None
    if f is None:
        out["note"] = ("Stt %s is '%s' in the sheet, the function list says '%s'" % (stt, funcs[stt]["name"], fn.get("description"))
                       if stt in funcs else "function not in the sheet (Stt %s, '%s')" % (stt, fn.get("description")))
        return out
    out.update(name=f["name"], flag=f["flag"])
    for i, (text, kind) in enumerate(f["items"], 1):
        roles, action, reply = split_transaction(text)
        out["items"].append({"id": "T%d" % i, "text": text, "roles": roles, "action": action, "reply": reply, "kind": kind})
    return out


def menu_path(fn):
    """The menu path of a function-list row as a list of labels ('A >> B >> C' -> ['A', 'B', 'C'])."""
    m = str(fn.get("menu") or "").strip()
    return [x.strip() for x in re.split(r"\s*>>\s*", m) if x.strip()] if m else []


def screen_names(fn):
    """Names the screen may carry in the code (its title): the row's name without the [# code] tag, and its description."""
    out = []
    for k in ("name", "description"):
        v = re.sub(r"^\s*\[[^\]]*\]\s*", "", str(fn.get(k) or "")).strip()
        if v and v not in out:
            out.append(v)
    return out


def find_siblings(profile, code, fn, trace=None, workspaces_root=None):
    """Other reports that share this one's screen: the same route, or the same main code files.

    Two reports on one screen usually mean one screen seen by two kinds of user (e.g. 1E_117 internal /
    1E_118 public company): a workbook that copies the sibling's cases then tests nothing new.
    Returns [{code, name, why: [...], workspace, delivered: [...]}].
    """
    import json
    out = {}
    trace = trace or {}
    top = [x["file"] for x in trace.get("top", [])[:3]]
    routes = set(trace.get("routes") or [])
    if (top or routes) and workspaces_root and os.path.isdir(workspaces_root):
        for d in os.listdir(workspaces_root):
            tp = os.path.join(workspaces_root, d, "trace.json")
            if d == code or not os.path.exists(tp):
                continue
            with open(tp, encoding="utf-8") as f:
                t = json.load(f)
            same_route = routes & set(t.get("routes") or [])
            if same_route:
                out.setdefault(d, {"code": d, "name": None, "why": []})["why"].append("cùng màn hình %s" % ", ".join(sorted(same_route)))
            same = len(set(top) & set(x["file"] for x in t.get("top", [])[:3]))
            if same >= 2:
                out.setdefault(d, {"code": d, "name": None, "why": []})["why"].append("trùng %d/3 file code đứng đầu trace" % same)
    if out:
        try:
            _, rows = _function_rows(profile)
        except SystemExit:
            rows = []
        for raw in rows:
            for c, s in out.items():
                if s["name"] is None and _has_code(raw, c, _code_columns(profile)):
                    s["name"] = _logical(profile, raw).get("name")
    for s in out.values():
        ws = os.path.join(workspaces_root or "", s["code"])
        s["workspace"] = ws if workspaces_root and os.path.isdir(ws) else None
        s["delivered"] = [os.path.basename(p) for p in find_previous(profile, s["code"])]
    return sorted(out.values(), key=lambda s: s["code"])


def find_standard(profile):
    """(path of the standard document or None, path of the profile's summary or None)."""
    st = profile.get("inputs.standard", {}) or {}
    docs = _tools(profile, st["glob"]) if st.get("glob") else []
    summ = profile.path(st["summary"]) if st.get("summary") else None
    return (docs[0] if docs else None), (summ if summ and os.path.exists(summ) else None)


_GROUP = re.compile(r"#\s*(\d[A-Z](?:\.[IVX]+(?:\.\d+)*)?)\s*_")
_CONV_ITEM = re.compile(r"^\|\s*(Q\d+)\s*\|\s*([^|]*)\|")


def group_key(fn):
    """'# 1D.I_NHÓM CHỨC NĂNG ...' -> '1D.I' (None when the function list gives no group, e.g. '#N/A')."""
    m = _GROUP.search(str((fn or {}).get("group") or ""))
    return m.group(1) if m else None


def group_conventions(profile, key):
    """Path of the group's conventions file (rules/groups/<key>.md, or profile rules.groups), whether or not it exists."""
    if not key:
        return None
    return os.path.join(profile.path(profile.get("rules.groups") or "rules/groups"), "%s.md" % key)


def convention_items(path):
    """{'Q1': 'decision text', ...} of a conventions file (table rows '| Qn | decision | ...')."""
    out = {}
    if path and os.path.exists(path):
        with open(path, encoding="utf-8") as f:
            for line in f:
                m = _CONV_ITEM.match(line.strip())
                if m:
                    out[m.group(1)] = " ".join(m.group(2).split())
    return out


def find_previous(profile, code):
    pat = profile.get("inputs.previous_workbooks")
    if not pat:
        return []
    pat = pat.replace("{code_nodash}", code.replace("_", "")).replace("{code}", code)
    return sorted(_tools(profile, pat), key=os.path.getmtime)
