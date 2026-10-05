"""Evidence cards for the 'Hình ảnh lỗi' sheet: one picture per bug (or per way a bug shows), rendered from HTML with
the kit's Chromium, so a reviewer sees the bug without re-running anything.

A card reads top to bottom like a bug report:
  header      BUG-01 (1/2) [Cao · Mở - còn lỗi 05/10] title
  meta        Tái hiện <date time> · Tài khoản · môi trường · bản build · Màn hình
  boxes       CÁC BƯỚC TÁI HIỆN | KẾT QUẢ THỰC TẾ (red) / KẾT QUẢ MONG ĐỢI (green)
  panels      ① screenshot (cropped, wrong element outlined)  ② exported Excel (wrong cells red, correct value in an
              extra blue column)  ③ source rows (row the standard picks green, row the report uses red)  ④ API text
  note        yellow box (same cause elsewhere, how many cells)
  footer      Căn cứ · Nguyên nhân (code)

Used from build_workbook.py:

    from reportkit import evidence as E
    run = ws.latest_run(); shots = run + "/shots/"
    meta = dict(account="gms", env="ids-dev-gms.net", build="frontend cập nhật 05/10/2026 11:30",
                screen="Thống kê > Thống kê CTĐC > BC quản trị từng công ty")
    cards = [
        E.card("BUG-01",                     # title / severity / status / basis / cause come from 'Danh sách lỗi'
               steps=["Tên công ty = **CTCP CTCBIO Việt Nam**; Từ năm = Đến năm = 2023", "Bấm **Xuất Excel**"],
               actual="Cột **Năm 2023** lấy bản **chờ duyệt** 1146790: STT 27 = 0 ...",
               expected="Lấy bản **đã đóng (CLOSE)** 1118103: STT 27 = 1 ...",
               panels=[E.run_shot(run, "U05", "Màn hình lúc xuất (bộ lọc + thông báo thành công)",
                                  mark=["toast"], show=["filters"]),      # positions recorded by rt check
                       E.excel(shots + "U05_R018_BaoCaoQuanTri.xlsx", "File Excel đã xuất - ô sai tô đỏ",
                               rows=["1-8", 27, 31], bad=["E27", "E31"],
                               extra={8: "Giá trị đúng (Tester ghi thêm)", 27: "1 (bản CLOSE 1118103)"}),
                       E.table("Dữ liệu nguồn - SELECT chỉ đọc lúc 05/10/2026 14:48",
                               ["COMPANY_DATA.ID", "Trạng thái", "STT 27", "Vai trò"],
                               [[1118103, "CLOSE", 1, "Quy chuẩn I.2 chọn bản này"],
                                [1146790, "REVIEWED", 0, "Báo cáo đang lấy bản này"]], roles=["ok", "bad"])],
               note="Cùng nguyên nhân: HTV 2022, CSM 2022. Toàn bộ 27 ô ở sheet 'Chi tiết sai lệch dữ liệu'.", **meta),
    ]
    wb.sheet_evidence(cards, title, note, no_image={"BUG-07": "lỗi thời gian chờ, không thể hiện trên hình"})

Text fields accept **bold** and line breaks. `captured` (dd/mm/yyyy HH:MM) defaults to the time of the first screenshot,
so a card never claims a date its picture was not taken on.

Cheap by design: E.run_shot crops / outlines from positions recorded by `rt check` (no need to open a screenshot), the
bug's wording comes from 'Danh sách lỗi', and render() lints every card (nothing marked, marked cell hidden, old
picture, too long) so only cards with warnings need to be looked at. E.text refuses credentials.
"""
import base64
import datetime
import html
import os
import re
import struct

CARD_W = 1360                      # px; the sheet shows the picture 1000 px wide
CIRCLED = "①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳"
WITHDRAWN = ("Rút", "Đã đóng")


# ------------------------------------------------------------------ building blocks
def card(bug, steps, actual, expected, panels, title="", status="", severity="", captured=None, account="", env="",
         build="", screen="", note="", basis="", cause=""):
    """One evidence picture. `panels` are E.shot / E.run_shot / E.excel / E.table / E.text, shown in order.
    title / severity / basis / cause left empty are taken from the bug's row of 'Danh sách lỗi' by sheet_evidence();
    status left empty becomes '<status of the row> - còn lỗi <dd/mm of captured>' (open bugs) - write only what differs."""
    if not re.match(r"^BUG-[\w-]+$", str(bug)):
        raise ValueError("card(%r): bug must be the id used in 'Danh sách lỗi' (BUG-01)" % bug)
    for k, v in (("actual", actual), ("expected", expected)):
        if not v:
            raise ValueError("card(%s): '%s' is required" % (bug, k))
    if not steps or isinstance(steps, str):
        raise ValueError("card(%s): 'steps' is a list of screen actions" % bug)
    if not panels:
        raise ValueError("card(%s): needs at least one panel (shot / excel / table / text)" % bug)
    if not captured:
        shot = next((p for p in panels if p["kind"] == "shot"), None)
        if not shot:
            raise ValueError("card(%s): 'captured' (dd/mm/yyyy HH:MM of the reproduction) is required when there is no screenshot" % bug)
        captured = datetime.datetime.fromtimestamp(os.path.getmtime(shot["path"])).strftime("%d/%m/%Y %H:%M")
    return dict(bug=bug, title=title, steps=list(steps), actual=actual, expected=expected, panels=list(panels),
                status=status, severity=severity, captured=captured, account=account, env=env, build=build,
                screen=screen, note=note, basis=basis, cause=cause, part="")


def shot(path, caption, crop=None, boxes=None):
    """A screenshot (PNG). crop=(x, y, w, h) in image pixels; boxes=[(x, y, w, h) or (x, y, w, h, 'label')] outlined in red
    (image pixels, before cropping). Outline the element itself at capture time with the UI step
    `screenshot: {name, selector, highlight}` when you can - it needs no coordinates."""
    if not os.path.exists(path):
        raise ValueError("shot: %s not found" % path)
    marked = 0
    if os.path.exists(path + ".json"):                     # written by the UI step screenshot: {highlight: ...}
        import json
        with open(path + ".json", encoding="utf-8") as f:
            marked = int(json.load(f).get("highlighted", 0))
    return dict(kind="shot", path=path, caption=caption, crop=crop, boxes=list(boxes or []), highlighted=marked)


def run_shot(run_dir, check_id, caption, mark=(), show=(), pad=40):
    """The final screenshot of a UI entry, cropped and outlined from the positions recorded during `rt check` -
    no need to open the picture to find coordinates.
      mark  names to outline in red: observation names of the entry, or its `boxes:` names; 'name#2' = the 3rd match only
      show  names to keep in the picture without outlining (e.g. the filter panel)
    The crop is the union of mark + show, plus `pad` px."""
    import json
    p = os.path.join(run_dir, "%s.json" % check_id)
    if not os.path.exists(p):
        raise ValueError("run_shot: %s not found" % p)
    with open(p, encoding="utf-8") as f:
        rec = (json.load(f).get("boxes") or {})
    known, path = rec.get("boxes") or {}, rec.get("shot")
    if not path or not os.path.exists(path):
        raise ValueError("run_shot(%s): no final screenshot with positions in this run (re-run the entry with the current kit)" % check_id)

    def pick(names):
        out = []
        for n in names:
            base, _, i = str(n).partition("#")
            if base not in known:
                raise ValueError("run_shot(%s): no position for %r; recorded: %s" % (check_id, base, sorted(known)))
            out += [known[base][int(i)]] if i else known[base]
        return out
    marks, shown_ = pick(mark), pick(show)
    crop = None
    if marks or shown_:
        w, h = png_size(path)
        allb = marks + shown_
        x0, y0 = max(0, min(b[0] for b in allb) - pad), max(0, min(b[1] for b in allb) - pad)
        x1, y1 = min(w, max(b[0] + b[2] for b in allb) + pad), min(h, max(b[1] + b[3] for b in allb) + pad)
        crop = (x0, y0, x1 - x0, y1 - y0)
    return shot(path, caption, crop=crop, boxes=marks)


def excel(path, caption, sheet=None, cells=None, rows=None, cols=None, bad=None, ok=None, extra=None):
    """A range of a downloaded Excel file, drawn from the file itself (values with their number format, fonts, fills,
    borders, merges, column widths) - so a format bug shows as it is in the file.
      cells  'A1:H60' (default: the used range)
      rows / cols  the rows / columns to show (others hidden, say so in the caption): [1, 2, '8-12', 27]; ['A-C', 'F']
      bad / ok     cells to mark red / green: ['E27', 'E31:E35']
      extra        {row: text} an extra blue column on the right, for the correct value (not in the file)
    """
    if not os.path.exists(path):
        raise ValueError("excel: %s not found" % path)
    return dict(kind="excel", path=path, caption=caption, sheet=sheet, cells=cells, rows=rows, cols=cols,
                bad=list(bad or []), ok=list(ok or []), extra=dict(extra or {}))


def table(caption, headers, rows, roles=None, bad=None):
    """Rows from `rt sql` / an API response. roles: per row 'ok' (green) / 'bad' (red) / None; bad: [(row, col)] 0-based cells."""
    if roles and len(roles) != len(rows):
        raise ValueError("table(%r): one role per row" % caption)
    return dict(kind="table", caption=caption, headers=list(headers), rows=[list(r) for r in rows],
                roles=list(roles or [None] * len(rows)), bad={tuple(x) for x in (bad or [])})


SECRET = re.compile(r"(authorization|password|passwd|mật khẩu|access_token|refresh_token|set-cookie|cookie)[\"']?\s*[:=]"
                    r"|bearer\s+[\w.-]{12,}", re.I)


def text(caption, body, max_chars=3000):
    """Verbatim text: a request and the start of its response, a log line, an error message. Refuses credentials."""
    body = str(body)
    m = SECRET.search(body)
    if m:
        raise ValueError("text(%r): looks like a credential (%r) - remove headers / tokens before putting it in a picture" % (caption, m.group(0)))
    if len(body) > max_chars:
        body = body[:max_chars] + "\n… (cắt bớt, tổng %d ký tự)" % len(body)
    return dict(kind="text", caption=caption, body=body)


# ------------------------------------------------------------------ HTML
def _md(s):
    """Escape, then **bold** and line breaks."""
    s = html.escape(str(s if s is not None else ""))
    return re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s).replace("\n", "<br>")


def png_size(path):
    with open(path, "rb") as f:
        head = f.read(24)
    if head[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError("%s is not a PNG" % path)
    return struct.unpack(">II", head[16:24])


def _data_uri(path):
    with open(path, "rb") as f:
        return "data:image/png;base64," + base64.b64encode(f.read()).decode()


def _shot_html(p):
    w, h = png_size(p["path"])
    x, y, cw, ch = p["crop"] or (0, 0, w, h)
    cw, ch = min(cw, w - x), min(ch, h - y)
    k = min(1.0, (CARD_W - 2) / float(cw))
    boxes = ""
    for b in p["boxes"]:
        bx, by, bw, bh = b[:4]
        label = ('<span class="blabel">%s</span>' % _md(b[4])) if len(b) > 4 else ""
        boxes += '<div class="box" style="left:%dpx;top:%dpx;width:%dpx;height:%dpx">%s</div>' % (
            (bx - x) * k, (by - y) * k, bw * k, bh * k, label)
    return ('<div class="shot" style="width:%dpx;height:%dpx"><img src="%s" style="left:%dpx;top:%dpx;width:%dpx;height:%dpx">%s</div>'
            % (cw * k, ch * k, _data_uri(p["path"]), -x * k, -y * k, w * k, h * k, boxes))


def _span(items, conv):
    out = []
    for it in items:
        if isinstance(it, str) and "-" in it:
            a, b = it.split("-", 1); out += list(range(conv(a), conv(b) + 1))
        else:
            out.append(conv(it))
    return out


def _date_fmt(v, f):
    f = re.sub(r"\[[^\]]*\]", "", f.split(";")[0]).replace("\\", "").replace('"', "")
    toks = re.findall(r"yyyy|yy|mm|m|dd|d|hh|h|ss|s|.", f, flags=re.I)
    out, after_h = [], False
    for i, t in enumerate(toks):
        tl = t.lower()
        if tl in ("yyyy", "yy"):
            out.append("%04d" % v.year if tl == "yyyy" else "%02d" % (v.year % 100))
        elif tl in ("mm", "m"):
            nxt = next((x.lower() for x in toks[i + 1:] if x.lower()[:1] in "dhsy"), "")
            n = getattr(v, "minute", 0) if (after_h or nxt.startswith("s")) else v.month
            out.append("%02d" % n if tl == "mm" else str(n))
        elif tl in ("dd", "d"):
            out.append("%02d" % v.day if tl == "dd" else str(v.day))
        elif tl in ("hh", "h"):
            out.append("%02d" % getattr(v, "hour", 0) if tl == "hh" else str(getattr(v, "hour", 0)))
        elif tl in ("ss", "s"):
            out.append("%02d" % getattr(v, "second", 0))
        else:
            out.append(t); continue
        after_h = tl in ("hh", "h")
    return "".join(out)


def shown(c):
    """The text Excel shows for a cell: its number / date format applied (common formats; en-US grouping 1,234.56)."""
    v, f = c.value, c.number_format or "General"
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (datetime.datetime, datetime.date)):
        return _date_fmt(v, f) if f != "General" else v.strftime("%d/%m/%Y")
    if isinstance(v, (int, float)):
        f0 = f.split(";")[0]
        if f0 in ("General", "@"):
            return ("%d" % v) if float(v).is_integer() else ("%.10g" % v)
        pct = "%" in f0
        m = re.search(r"\.(0+)", f0)
        dec = len(m.group(1)) if m else 0
        s = ("{:,.%df}" if "," in f0 else "{:.%df}") % dec
        return s.format(v * 100 if pct else v) + ("%" if pct else "")
    return str(v)


def _argb(c):
    rgb = getattr(c, "rgb", None) if c is not None else None
    return ("#" + rgb[-6:]) if isinstance(rgb, str) and len(rgb) >= 6 and c.type == "rgb" else None


def _cell_css(c, w_px):
    css = []
    f = c.font
    if f is not None:
        if f.b: css.append("font-weight:bold")
        if f.i: css.append("font-style:italic")
        if f.u: css.append("text-decoration:underline")
        if f.sz: css.append("font-size:%.1fpx" % (float(f.sz) * 4 / 3))
        if f.name: css.append("font-family:'%s',serif" % f.name.replace("'", ""))
        col = _argb(f.color)
        if col: css.append("color:" + col)
    if c.fill is not None and c.fill.fill_type == "solid":
        col = _argb(c.fill.fgColor)
        if col: css.append("background:" + col)
    al = c.alignment
    h = al.horizontal if al is not None else None
    if h in ("left", "center", "right"):
        css.append("text-align:" + h)
    elif isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
        css.append("text-align:right")
    css.append("vertical-align:%s" % {"center": "middle", "top": "top"}.get(al.vertical if al is not None else None, "bottom"))
    css.append("white-space:%s" % ("pre-wrap" if (al is not None and al.wrap_text) else "pre"))
    b = c.border
    for side in ("left", "right", "top", "bottom"):
        s = getattr(b, side, None) if b is not None else None
        if s is not None and s.style:
            css.append("border-%s:%dpx solid %s" % (side, 2 if s.style in ("medium", "thick", "double") else 1, _argb(s.color) or "#000"))
    css.append("max-width:%dpx" % w_px)
    return ";".join(css)


def _excel_html(p):
    import openpyxl
    from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl, range_boundaries
    wb = openpyxl.load_workbook(p["path"], data_only=True)
    sh = wb[p["sheet"]] if p["sheet"] else wb.worksheets[0]
    c1, r1, c2, r2 = range_boundaries(p["cells"] or sh.dimensions)
    rows = [r for r in (_span(p["rows"], int) if p["rows"] else range(r1, r2 + 1)) if r1 <= r <= r2]
    cols = [c for c in (_span(p["cols"], lambda x: ci(x) if isinstance(x, str) else int(x)) if p["cols"] else range(c1, c2 + 1)) if c1 <= c <= c2]
    mark = {}
    for kind in ("bad", "ok"):
        for ref in p[kind]:
            a, b, c, d = range_boundaries(ref if ":" in ref else "%s:%s" % (ref, ref))
            for rr in range(b, d + 1):
                for cc in range(a, c + 1):
                    mark[(rr, cc)] = kind
    vis_r, vis_c, span, covered = set(rows), set(cols), {}, set()
    hidden = ["%s%d" % (gl(c), r) for (r, c), k in sorted(mark.items()) if k == "bad" and (r not in vis_r or c not in vis_c)]
    p["_warn"] = (["ô đánh dấu sai bị ẩn: " + ", ".join(hidden)] if hidden else []) + \
                 ["giá trị đúng ở dòng bị ẩn: %s" % r for r in p["extra"] if r not in vis_r]
    for m in sh.merged_cells.ranges:
        rs = [r for r in range(m.min_row, m.max_row + 1) if r in vis_r]
        cs = [c for c in range(m.min_col, m.max_col + 1) if c in vis_c]
        if not rs or not cs:
            continue
        span[(rs[0], cs[0])] = (len(rs), len(cs))
        covered |= {(r, c) for r in rs for c in cs} - {(rs[0], cs[0])}
    widths = {c: int((sh.column_dimensions[gl(c)].width or 8.43) * 7 + 5) for c in cols}
    out = ['<table class="xl"><colgroup>%s%s</colgroup>' % ("".join('<col style="width:%dpx">' % widths[c] for c in cols),
                                                             '<col style="width:260px">' if p["extra"] else "")]
    for r in rows:
        ht = sh.row_dimensions[r].height
        out.append('<tr style="height:%dpx">' % (ht * 4 / 3 if ht else 20))
        for c in cols:
            if (r, c) in covered:
                continue
            cell = sh.cell(r, c)
            rs, cs = span.get((r, c), (1, 1))
            css = _cell_css(cell, sum(widths[x] for x in cols[cols.index(c):cols.index(c) + cs]))
            cls = {"bad": ' class="bad"', "ok": ' class="ok"'}.get(mark.get((r, c)), "")
            out.append('<td%s%s style="%s">%s</td>' % (cls, (' rowspan="%d" colspan="%d"' % (rs, cs)) if (rs, cs) != (1, 1) else "",
                                                      css, html.escape(shown(cell))))
        if p["extra"]:
            t = p["extra"].get(r)
            out.append('<td class="extra">%s</td>' % _md(t) if t else '<td class="noextra"></td>')
        out.append("</tr>")
    out.append("</table>")
    return "".join(out)


def _table_html(p):
    out = ['<table class="src"><tr>%s</tr>' % "".join("<th>%s</th>" % _md(h) for h in p["headers"])]
    for i, (row, role) in enumerate(zip(p["rows"], p["roles"])):
        out.append('<tr class="%s">%s</tr>' % (role or "", "".join(
            '<td%s>%s</td>' % (' class="badcell"' if (i, j) in p["bad"] else "", _md(v)) for j, v in enumerate(row))))
    out.append("</table>")
    return "".join(out)


CSS = """
*{box-sizing:border-box} body{margin:0;background:#fff;font-family:'Segoe UI',Arial,sans-serif;color:#101828}
.card{width:%(w)dpx;padding:20px 20px 16px}
.hdr{display:flex;align-items:center;gap:14px;border-bottom:2px solid #B42318;padding-bottom:8px}
.bug{font-size:27px;font-weight:700;color:#B42318;white-space:nowrap}
.chip{background:#FEE4E2;color:#B42318;border-radius:12px;padding:2px 12px;font-size:14px;white-space:nowrap}
.chip.grey{background:#EAECF0;color:#475467} .bug.grey{color:#475467} .hdr.grey{border-color:#98A2B3}
.ttl{font-size:21px}
.meta{font-size:13px;color:#475467;margin:8px 0 12px}
.grid{display:grid;grid-template-columns:1fr 1fr;gap:14px;margin-bottom:14px}
.boxx{border:1px solid #D0D5DD;border-radius:6px;padding:10px 14px;font-size:14.5px;line-height:1.5}
.boxx h4{margin:0 0 6px;font-size:13px;color:#344054;letter-spacing:.3px}
.boxx ol{margin:0;padding-left:20px}
.act{background:#FEF3F2;border-color:#FDA29B} .act h4{color:#B42318}
.exp{background:#ECFDF3;border-color:#75E0A7;margin-top:12px} .exp h4{color:#067647}
.cap{font-size:14.5px;margin:12px 0 6px}
.shot{position:relative;overflow:hidden;border:1px solid #D0D5DD}
.shot img{position:absolute;max-width:none}
.box{position:absolute;border:3px solid #D92D20;border-radius:3px}
.blabel{position:absolute;left:-3px;top:-24px;background:#D92D20;color:#fff;font-size:12px;padding:1px 6px;white-space:nowrap}
table.xl{border-collapse:collapse;table-layout:fixed;background:#fff;font-size:16px}
table.xl td{padding:1px 4px;overflow:hidden;text-overflow:clip}
table.xl td.bad{background:#FFC7CE!important;color:#C00000!important;font-weight:bold}
table.xl td.ok{background:#C6EFCE!important}
table.xl td.extra{background:#DDEBF7;color:#1F4E79;font-style:italic;border:1px solid #2F5597;white-space:pre-wrap}
.xlwrap{border:1px solid #D0D5DD;padding:8px;display:inline-block;max-width:100%%;overflow:hidden}
table.src{border-collapse:collapse;font-size:13px;width:100%%}
table.src th{background:#F2F4F7;text-align:left;border:1px solid #D0D5DD;padding:4px 6px}
table.src td{border:1px solid #D0D5DD;padding:4px 6px;vertical-align:top}
table.src tr.ok td{background:#ECFDF3} table.src tr.bad td{background:#FEF3F2} td.badcell{color:#C00000;font-weight:bold}
pre{background:#F9FAFB;border:1px solid #D0D5DD;padding:8px 10px;font-size:12.5px;white-space:pre-wrap;word-break:break-all;margin:0}
.note{background:#FFFAEB;border:1px solid #FEDF89;border-radius:6px;padding:8px 12px;font-size:13.5px;margin-top:14px}
.foot{font-size:12.5px;color:#475467;margin-top:12px;border-top:1px solid #EAECF0;padding-top:8px;line-height:1.5}
.foot b{color:#344054}
""" % {"w": CARD_W}


def card_html(c):
    grey = " grey" if str(c["status"]).startswith(WITHDRAWN) else ""
    head = "%s%s" % (c["bug"], (" (%s)" % c["part"]) if c["part"] else "")
    chip = " · ".join(x for x in (c["severity"], c["status"]) if x)
    meta = " · ".join(x for x in ("Tái hiện: <b>%s</b>" % _md(c["captured"]),
                                  ("Tài khoản: <b>%s</b>" % _md(c["account"])) if c["account"] else "",
                                  _md(c["env"]) if c["env"] else "", _md(c["build"]) if c["build"] else "",
                                  ("Màn hình: %s" % _md(c["screen"])) if c["screen"] else "") if x)
    parts = ['<div class="card"><div class="hdr%s"><span class="bug%s">%s</span>%s<span class="ttl">%s</span></div>'
             % (grey, grey, _md(head), ('<span class="chip%s">%s</span>' % (grey, _md(chip))) if chip else "", _md(c["title"])),
             '<div class="meta">%s</div>' % meta,
             '<div class="grid"><div class="boxx"><h4>CÁC BƯỚC TÁI HIỆN</h4><ol>%s</ol></div>' % "".join("<li>%s</li>" % _md(s) for s in c["steps"]),
             '<div><div class="boxx act"><h4>KẾT QUẢ THỰC TẾ</h4>%s</div>' % _md(c["actual"]),
             '<div class="boxx exp"><h4>KẾT QUẢ MONG ĐỢI</h4>%s</div></div></div>' % _md(c["expected"])]
    for i, p in enumerate(c["panels"]):
        parts.append('<div class="cap">%s %s</div>' % (CIRCLED[i] if i < len(CIRCLED) else "(%d)" % (i + 1), _md(p["caption"])))
        if p["kind"] == "shot":
            parts.append(_shot_html(p))
        elif p["kind"] == "excel":
            parts.append('<div class="xlwrap">%s</div>' % _excel_html(p))
        elif p["kind"] == "table":
            parts.append(_table_html(p))
        else:
            parts.append("<pre>%s</pre>" % html.escape(p["body"]))
    if c["note"]:
        parts.append('<div class="note">%s</div>' % _md(c["note"]))
    foot = " · ".join(x for x in (("<b>Căn cứ:</b> %s" % _md(c["basis"])) if c["basis"] else "",
                                  ("<b>Nguyên nhân (code):</b> %s" % _md(c["cause"])) if c["cause"] else "") if x)
    if foot:
        parts.append('<div class="foot">%s</div>' % foot)
    parts.append("</div>")
    return '<!doctype html><html><head><meta charset="utf-8"><style>%s</style></head><body>%s</body></html>' % (CSS, "".join(parts))


# ------------------------------------------------------------------ render
def number(cards):
    """Sets 'part' = 'i/n' on bugs that have several cards. Cards of one bug must be consecutive."""
    seen, order = {}, []
    for c in cards:
        if c["bug"] in seen and order[-1] != c["bug"]:
            raise ValueError("cards of %s are not consecutive" % c["bug"])
        seen.setdefault(c["bug"], []).append(c)
        if not order or order[-1] != c["bug"]:
            order.append(c["bug"])
    for group in seen.values():
        for i, c in enumerate(group, 1):
            c["part"] = "%d/%d" % (i, len(group)) if len(group) > 1 else ""
    return cards


def fill_from_bugs(cards, bugs):
    """Fills the empty title / severity / status / basis / cause of each card from its bug row.
    bugs: {bug id: {'severity', 'status', 'title', 'basis', 'cause'}} (sheet_evidence reads them from 'Danh sách lỗi')."""
    for c in cards:
        row = bugs.get(c["bug"]) or {}
        for k in ("title", "severity", "basis", "cause"):
            if not c[k]:
                c[k] = row.get(k) or ""
        if not c["status"]:
            st = row.get("status") or ""
            c["status"] = ("%s - còn lỗi %s" % (st, c["captured"][:5])) if st in ("Mới", "Mở") else st
    return cards


def lint(c, today=None):
    """What a reviewer would reject, checked without looking at the picture. [] = nothing found."""
    out, marked = [], False
    for p in c["panels"]:
        if p["kind"] == "shot":
            marked = marked or bool(p["boxes"]) or p.get("highlighted", 0) > 0
            if p["crop"]:
                w, h = png_size(p["path"])
                x, y, cw, ch = p["crop"]
                if cw <= 0 or ch <= 0 or x >= w or y >= h:
                    out.append("vùng cắt nằm ngoài ảnh %s" % os.path.basename(p["path"]))
        elif p["kind"] == "excel":
            marked = marked or bool(p["bad"]) or bool(p["extra"])
            out += p.get("_warn", [])
        elif p["kind"] == "table":
            marked = marked or any(r == "bad" for r in p["roles"]) or bool(p["bad"])
        else:
            marked = True                              # an error text is its own mark
    if not marked:
        out.append("không có chỗ nào được đánh dấu (boxes / highlight / bad / roles)")
    today = today or datetime.date.today().strftime("%d/%m/%Y")
    if not str(c["captured"]).startswith(today) and "chưa kiểm lại" not in str(c["status"]):
        out.append("ảnh chụp %s, không phải hôm nay: trạng thái phải ghi 'chưa kiểm lại'" % c["captured"])
    if not c["title"] or not c["status"]:
        out.append("thiếu tên lỗi / trạng thái")
    return out


def render(cards, out_dir):
    """Renders every card to <out_dir>/<BUG-01_1>.png (+ .html for inspection), sets card['png'] and card['warnings'],
    and prints one line per card. Open a card only when it has warnings (and the first one of a run, for the layout)."""
    from playwright.sync_api import sync_playwright
    number(cards)
    os.makedirs(out_dir, exist_ok=True)
    paths = []
    with sync_playwright() as pw:
        browser = pw.chromium.launch()
        pg = browser.new_page(viewport={"width": CARD_W + 40, "height": 900}, device_scale_factor=1)
        for c in cards:
            base = os.path.join(out_dir, "%s_%s" % (c["bug"], (c["part"] or "1").split("/")[0]))
            doc = card_html(c)
            with open(base + ".html", "w", encoding="utf-8") as f:
                f.write(doc)
            pg.set_content(doc, wait_until="load")
            pg.locator(".card").screenshot(path=base + ".png")
            c["png"] = base + ".png"; paths.append(c["png"])
            w, h = png_size(c["png"])
            c["warnings"] = lint(c) + (["thẻ dài %d px: tách thành nhiều hình (1/2), (2/2)" % h] if h > 3500 else [])
            print("evidence %s%s: %dx%d px, %d panels, %s" % (c["bug"], (" (%s)" % c["part"]) if c["part"] else "", w, h,
                                                            len(c["panels"]), "; ".join(c["warnings"]) or "OK"))
        browser.close()
    return paths
