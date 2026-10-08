import json
import os

import openpyxl
import pytest

from reportkit import inputs, workbook as W
from test_core import _analysis, _mini_profile

FL = {"function_list": {"glob": "list.xlsx", "sheet": "List chức năng", "header_contains": "Mã Jira",
                        "code_columns": ["New UC Name"],
                        "columns": {"stt": "Stt", "name": "New UC Name", "description": "Mô tả yêu cầu của người dùng",
                                    "system": "Phân hệ"}},
      "transactions": {"sheet": "Chi tiết transaction", "system": "IDS",
                       "columns": {"stt": "Stt", "text": "Mô tả yêu cầu", "kind": "Phân loại dữ liệu"}}}


def _profile(tmp_path):
    prof = _mini_profile(tmp_path, extra={"inputs": FL})
    wb = openpyxl.Workbook(); sh = wb.active; sh.title = "List chức năng"
    sh.append(["KBKT bàn giao"])
    sh.append(["Stt", "Phân hệ", "Mô tả yêu cầu của người dùng", "Mã Jira", "New UC Name 11.06"])
    sh.append(["E", "IDS", "NHÓM IN ẤN"])
    sh.append([117, "IDS", "Trang tổng hợp công bố thông tin nội bộ", "J-1", "[# 1E_117] Trang tổng hợp CBTT nội bộ"])
    sh.append([122, "IDS", "Tra cứu thông tin báo cáo tài chính", "J-2", "[# 1E_122] Tra cứu BCTC"])
    sh.append([123, "IDS", "Tên đã đổi trong danh sách", "J-3", "[# 1E_123] Tra cứu khác"])
    sh.append([117, "KT", "Báo cáo của phân hệ kiểm toán", "J-4", "[# 2E_118] Báo cáo KT"])
    tx = wb.create_sheet("Chi tiết transaction")
    tx.append(["Stt", "Mô tả yêu cầu", "Phân loại dữ liệu"])
    tx.append(["E", "NHÓM KHAI THÁC DỮ LIỆU"])
    tx.append([117, "Trang tổng hợp công bố thông tin nội bộ", "Sửa"])
    tx.append([None, "LĐGSĐC; CVGSĐC; LĐCBPH, CVCBPH xem tổng hợp báo cáo giao dịch. Hệ thống hiển thị thông tin thống kê",
               "Dữ liệu đầu vào"])
    tx.append([None, "Hệ thống chạy job tổng hợp. Hệ thống tính lại số liệu", "Yêu cầu truy vấn"])
    tx.append([None, "CVGSĐC Nhập   tiêu chí", None])
    tx.append([122, "Tra cứu thông tin báo cáo tài chính", None])
    tx.append([None, "LĐGSĐC xuất Excel. Hệ thống tạo file theo mẫu", "Dữ liệu đầu ra"])
    tx.append([123, "Tên cũ của chức năng", None])
    tx.append([None, "LĐGSĐC xem", "Dữ liệu đầu vào"])
    tx.append(["F", "TÍCH HỢP"])
    tx.append([None, "dòng lạc sau tiêu đề nhóm, không thuộc chức năng nào", None])
    wb.save(os.path.join(prof.tools_dir, "list.xlsx"))
    return prof


def test_transactions_joined_on_stt_with_roles_and_flag(tmp_path):
    prof = _profile(tmp_path)
    tx = inputs.find_transactions(prof, inputs.find_function(prof, "1E_117"))
    assert tx["flag"] == "Sửa" and tx["note"] is None and [t["id"] for t in tx["items"]] == ["T1", "T2", "T3"]
    t1, t2, t3 = tx["items"]
    assert t1["roles"] == ["LĐGSĐC", "CVGSĐC", "LĐCBPH", "CVCBPH"] and t1["kind"] == "Dữ liệu đầu vào"
    assert t1["action"] == "xem tổng hợp báo cáo giao dịch" and t1["reply"] == "Hệ thống hiển thị thông tin thống kê"
    assert t2["roles"] == [] and t2["action"] == "Hệ thống chạy job tổng hợp" and t2["reply"] == "Hệ thống tính lại số liệu"
    assert t3["roles"] == ["CVGSĐC"] and t3["action"] == "Nhập tiêu chí" and t3["kind"] is None


def test_transactions_not_listed(tmp_path):
    prof = _profile(tmp_path)
    kt = inputs.find_transactions(prof, inputs.find_function(prof, "2E_118"))
    assert kt["items"] == [] and "only IDS" in kt["note"]          # same Stt 117, other system: never joined
    renamed = inputs.find_transactions(prof, inputs.find_function(prof, "1E_123"))
    assert renamed["items"] == [] and "Tên cũ" in renamed["note"]  # Stt matches, name does not: not trusted
    (tmp_path / "p2").mkdir()
    assert inputs.find_transactions(_mini_profile(tmp_path / "p2"), {"stt": 117}) is None   # no sheet in the profile


def test_transactions_are_coverage_requirements(tmp_path):
    prof = _profile(tmp_path); ws = prof.workspace("1E_122"); _analysis(ws)
    tx = inputs.find_transactions(prof, inputs.find_function(prof, "1E_122"))
    with open(ws.p("inputs", "transactions.json"), "w", encoding="utf-8") as f:
        json.dump(tx, f, ensure_ascii=False)
    assert W.requirements(ws.dir, {})["T1"] == ("transaction", "LĐGSĐC xuất Excel. Hệ thống tạo file theo mẫu")

    def build(version, covers):
        wb = W.Workbook(prof, ws, name="n", screen="s", ticket="T", ascii_name="N", version=version)
        wb.chapter("Chức năng 1"); wb.cat("Chức năng")
        wb.tc("Kiểm tra xuất Excel", "1. Bấm [Xuất Excel]", "Có file theo mẫu", status="P", actual="Đạt. Có file.",
              covers=covers)
        wb.finish(run_date="07/10/2026", run_note="n")
        return wb.save()

    with pytest.raises(SystemExit) as e:
        build(1, ["R1"])
    assert "T1 (transaction)" in str(e.value)
    assert build(1, ["R1", "T1"]).endswith(".xlsx")
