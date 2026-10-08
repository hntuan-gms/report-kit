import datetime as dt
import os

import openpyxl
import pytest
import yaml

from reportkit import excel, locate, params, profile as P, safety
from reportkit.checks import assertions, data


# ---------------------------------------------------------------- params
def test_params_dates_and_arithmetic():
    v = params.base_vars(dt.date(2026, 1, 15))
    assert v["prev_month"] == 12 and v["prev_month_year"] == 2025
    assert params.render("{current_year-4} - {current_year}", v) == "2022 - 2026"
    assert params.render("{today:%Y-%m-%d}", v) == "2026-01-15"
    assert params.render("{current_year}", v) == 2026                 # lone numeric token -> int
    assert params.render("keep {unknown} and {fn(x)}", v) == "keep {unknown} and {fn(x)}"


def test_matrix_expands_cartesian():
    v = params.base_vars(dt.date(2026, 9, 28))
    combos = params.expand_matrix({"year": ["{current_year}", "{current_year-1}"], "lang": ["vi", "en"]}, v)
    assert len(combos) == 4 and {"year": 2025, "lang": "en"} in combos


def test_matrix_ranges_cover_every_value():
    v = params.base_vars(dt.date(2026, 9, 28))
    assert [c["year"] for c in params.expand_matrix({"year": "{current_year}..2022"}, v)] == [2026, 2025, 2024, 2023, 2022]
    assert [c["year"] for c in params.expand_matrix({"year": "2024..{current_year}"}, v)] == [2024, 2025, 2026]
    assert [c["y"] for c in params.expand_matrix({"y": ["2016..2017", "{current_year}"]}, v)] == [2016, 2017, 2026]
    assert params.expand_matrix({"rank": ["A", "B"]}, v) == [{"rank": "A"}, {"rank": "B"}]


# ---------------------------------------------------------------- safety
@pytest.mark.parametrize("sql", ["SELECT 1 FROM DUAL", "with x as (select 1 a from dual) select * from x",
                                 "SELECT 'DELETE me' TXT FROM DUAL", "SELECT DELETE_FLG, BEGIN_DATE FROM T"])
def test_sql_allowed(sql):
    safety.check_sql(sql)


@pytest.mark.parametrize("sql", ["DELETE FROM T", "UPDATE T SET A=1", "SELECT 1 FROM DUAL; DROP TABLE T",
                                 "BEGIN NULL; END;", "select * from t for update; commit",
                                 "/* SELECT */ INSERT INTO T VALUES (1)"])
def test_sql_refused(sql):
    with pytest.raises(safety.Refused):
        safety.check_sql(sql)


# ---------------------------------------------------------------- excel display
def test_shown_applies_formats():
    wb = openpyxl.Workbook(); ws = wb.active
    ws["A1"] = 2.0; ws["A1"].number_format = "#,##0"
    ws["A2"] = 1234.5; ws["A2"].number_format = "#,##0.00"
    ws["A3"] = dt.datetime(2026, 9, 28); ws["A3"].number_format = "dd/mm/yyyy"
    ws["A4"] = 0.1234; ws["A4"].number_format = "0.00%"
    assert [excel.shown(ws[c]) for c in ("A1", "A2", "A3", "A4")] == ["2", "1,234.50", "28/09/2026", "12.34%"]
    assert excel.std_number(25604) == "25,604" and excel.std_number(61.42) == "61.42"


# ---------------------------------------------------------------- preconditions
def test_preconditions():
    assert assertions.check("Hạng A: 57/2325", {"regex": "^Hạng A"}) == []
    assert assertions.check("Hạng B: 335/2325", {"regex": "^Hạng A"})
    assert assertions.check(4, {"min": 4}) == [] and assertions.check([], {"not_empty": True})
    assert assertions.check_all({"n": 3}, {"n": {"min": 4}, "missing": {"min": 1}}).keys() == {"n", "missing"}


# ---------------------------------------------------------------- data comparison helpers
def test_compare_numeric_and_text():
    assert data._same("2,433", 2433, {}) and data._same(61.42, "61.42", {})
    assert not data._same(2433, 2436, {}) and data._same(2433, 2436, {"tolerance": 5})
    assert data._same(" Hạng  A ", "Hạng A", {"numeric": False})
    k = data._keyed([{"thang": 1, "a": 5}, {"thang": 2, "a": 7}], ["thang"], {"DK": "a"})
    assert k == {("1",): {"DK": 5}, ("2",): {"DK": 7}}


# ---------------------------------------------------------------- locate
def test_fingerprints_from_free_text():
    txt = ("Lấy tin từ COMPANY_DATA với NEWS_TYPE_CD in ('DINH_KY','BAO_CAO'); trạng thái 'SUBMITTED_LATE'. "
           "Tên biểu đồ: XU HƯỚNG VI PHẠM CÔNG BỐ THÔNG TIN QUA CÁC NĂM")
    ids, phrases = locate.fingerprints(txt)
    assert {"COMPANY_DATA", "NEWS_TYPE_CD", "DINH_KY", "BAO_CAO", "SUBMITTED_LATE"} <= set(ids)
    assert any("XU HƯỚNG VI PHẠM" in p for p in phrases)


def test_labels_and_globs():
    assert locate.label_variants("CTĐC đăng ký với UBCKNN/ Danh sách (mới)")[:3] == [
        "ctđc đăng ký với ubcknn/ danh sách (mới)", "ctđc đăng ký với ubcknn", "danh sách (mới)"]
    assert locate._match("ids-frontend/src/app/app.routes.ts", "*-frontend/src/app/**/*.ts")          # '**/' = no folder too
    assert locate._match("ids-frontend/src/app/pages/x/a.component.ts", "*-frontend/src/app/**/*.ts")
    assert not locate._match("ids-frontend/src/assets/a.ts", "*-frontend/src/app/**/*.ts")


def test_menu_entries_keep_parents():
    text = """export const MENU = [
      { id: 'stats', title: 'menu.statistics', children: [
          { id: 'ctdc', title: 'menu.publicCompanyStats', children: [
              { id: 'gov', title: 'menu.governance', routerLink: '/statistics/ctdc/governance' },   // a comment with { brace
          ] },
      ] },
      { id: 'home', title: 'menu.dashboard', routerLink: '/dashboard' },
    ];"""
    e = {x["title"]: x for x in locate.menu_entries(text)}
    assert e["menu.governance"]["link"] == "/statistics/ctdc/governance"
    assert e["menu.governance"]["parents"] == ["menu.statistics", "menu.publicCompanyStats"]
    assert e["menu.dashboard"]["parents"] == [] and e["menu.statistics"]["link"] is None


def _fake_repo(root):
    files = {
        "web/src/assets/i18n/vi/menu.json": '{"menu": {"statistics": "Thống kê", "ctdc": "Thống kê CTĐC", "gov": "BC quản trị", "other": "BC quản trị"}}',
        "web/src/app/shared/menu.data.ts": """export const M = [
            { title: 'menu.statistics', children: [ { title: 'menu.ctdc', children: [
                { title: 'menu.gov', routerLink: '/statistics/ctdc/governance' } ] } ] },
            { title: 'menu.other', routerLink: '/elsewhere' } ];""",
        "web/src/app/pages/statistics/statistics.routes.ts": """export const routes = [
            { path: 'ctdc/governance', loadComponent: () => import('./gov/gov.component').then((c) => c.GovComponent) },
            { path: 'ctdc/other', loadComponent: () => import('./other/other.component').then((c) => c.OtherComponent) } ];""",
        "web/src/app/pages/statistics/gov/gov.component.ts": """import { ReportService } from '../../../service/report.service';
            export class GovComponent { private readonly reportService = inject(ReportService);
              readonly reportCode = ReportCode.R017_QUAN_TRI;
              load() { this.reportService.export(this.reportCode); } }""",
        "web/src/app/pages/statistics/gov/gov.component.html": "<h1>{{ 'statistics.gov.title' | translate }}</h1>",
        "web/src/app/service/report.service.ts": """export class ReportService { apiUrl = 'report';
              export(code) { return this.http.post(`${this.baseUrl}/${this.apiUrl}/${code}/export`, {}); }
              other() { return this.http.get('/unused/thing'); } }""",
        "api/src/main/java/R017QuanTriReport.java": """class R017QuanTriReport { ReportCode code() { return ReportCode.R017_QUAN_TRI; }
              String sql = "SELECT * FROM V_RPT_R017_QUAN_TRI"; String tpl = "templates/excel/R017_quan_tri.xlsx"; }""",
        "api/src/main/java/ReportCode.java": "enum ReportCode { " + ", ".join("R%03d_X" % i for i in range(1, 30)) + ", R017_QUAN_TRI }",
        "api/src/main/java/ReportController.java": '@RequestMapping("/{reportCode}") class ReportController { @PostMapping("/export") void e() {} }',
        "api/src/main/resources/db/R017_v.sql": "CREATE OR REPLACE VIEW V_RPT_R017_QUAN_TRI AS SELECT * FROM COMPANY_DATA JOIN FORMS ON 1=1",
        "api/src/main/resources/templates/excel/R017_quan_tri.xlsx": "",
    }
    for rel, body in files.items():
        p = root / rel; p.parent.mkdir(parents=True, exist_ok=True); p.write_text(body, encoding="utf-8")


def test_locate_follows_menu_to_backend(tmp_path):
    root = tmp_path / "repo"; _fake_repo(root)
    prof = _mini_profile(tmp_path, {"paths": {"tools_dir": str(tmp_path / "tools"), "work_root": str(tmp_path / "w"), "code_root": str(root)},
                                    "codemap": {"include": ["web/src/**/*.ts", "web/src/**/*.html", "web/src/**/*.json",
                                                            "api/src/**/*.java", "api/src/**/*.sql"],
                                                "i18n": "web/src/assets/i18n/vi/**/*.json", "frontend": {"a": "web"}}})
    ws = prof.workspace("X_1")
    fn = {"menu": "Thống kê >> Thống kê CTĐC >> BC quản trị >> Tab chi tiết", "name": "[# X_1] Báo cáo quản trị", "system_key": "a"}
    t = locate.locate(prof, ws, fn)
    assert t["method"] == "menu path" and t["routes"] == ["/statistics/ctdc/governance"]       # 'other' has no matching parents
    assert t["menu"][0]["unmatched_tail"] == ["Tab chi tiết"]
    assert t["components"] == ["web/src/app/pages/statistics/gov/gov.component.ts"]
    assert t["endpoints"] == ["POST /report/{}/export (report.service.ts.export)"]
    assert t["report_codes"] == ["R017"]
    be = [x["file"] for x in t["top"] if x["score"] is not None]
    assert be[0] == "api/src/main/java/R017QuanTriReport.java"                                   # the registry enum ranks below
    assert be.index("api/src/main/java/ReportCode.java") > 0
    assert {"V_RPT_R017_QUAN_TRI", "COMPANY_DATA", "FORMS"} <= set(t["sql_objects"])
    assert t["views"][0]["file"] == "api/src/main/resources/db/R017_v.sql"
    assert t["templates"] == ["api/src/main/resources/templates/excel/R017_quan_tri.xlsx"]
    assert locate.locate(prof, ws, fn).get("cached")
    md = locate.trace_md(t)
    assert "Tab chi tiết" in md and "R017QuanTriReport.java" in md


def test_locate_with_route_and_without_menu(tmp_path):
    root = tmp_path / "repo"; _fake_repo(root)
    prof = _mini_profile(tmp_path, {"paths": {"tools_dir": str(tmp_path / "tools"), "work_root": str(tmp_path / "w"), "code_root": str(root)},
                                    "codemap": {"include": ["web/src/**/*.ts", "web/src/**/*.json", "api/src/**/*.java"],
                                                "i18n": "web/src/assets/i18n/vi/**/*.json"}})
    t = locate.locate(prof, prof.workspace("X_2"), {"name": "x"}, route="/statistics/ctdc/governance")
    assert t["method"] == "route given" and t["components"][0].endswith("gov.component.ts")
    t = locate.locate(prof, prof.workspace("X_3"), {"name": "[# X_3] Không có trong code"})
    assert t["method"].startswith("name only") and "screen component" in t["not_found"]


# ---------------------------------------------------------------- profile + workspace
def test_profile_loading_and_codes(tmp_path, monkeypatch):
    d = tmp_path / "proj" / ".report-kit"; d.mkdir(parents=True)
    (d / "project.yaml").write_text(yaml.safe_dump({"name": "demo", "paths": {"tools_dir": "{home}/x", "work_root": str(tmp_path / "w")},
                                                   "systems": {"a": {"match": ["App A"], "web": "http://a", "api": "http://a/api"}}}),
                                    encoding="utf-8")
    (d / "project.local.yaml").write_text(yaml.safe_dump({"paths": {"tools_dir": str(tmp_path / "tools")}}), encoding="utf-8")
    monkeypatch.chdir(tmp_path / "proj")
    monkeypatch.setenv("RK_LOGIN_USER", "u1")
    prof = P.load()
    assert prof.name == "demo" and prof.tools_dir == str(tmp_path / "tools")
    assert prof.system_for_label("App A") == "a" and prof.secret("login_user") == "u1"
    with pytest.raises(SystemExit):
        prof.secret("db_password")
    assert P.norm_code("[# 1E_117]") == "1E_117" and P.norm_code("1e117") == "1E_117" and P.norm_code("2E_50") == "2E_50"
    ws = prof.workspace("1E117")
    ws.mark("start", x=1)
    assert ws.done("start") and os.path.isdir(ws.dir) and ws.latest_run() is None


def test_engine_rejects_duplicate_ids(tmp_path, monkeypatch):
    from reportkit.checks import engine
    d = tmp_path / ".report-kit"; d.mkdir()
    (d / "project.yaml").write_text(yaml.safe_dump({"name": "demo", "paths": {"work_root": str(tmp_path / "w")},
                                                   "systems": {"a": {"web": "http://a", "api": "http://a/api"}}}), encoding="utf-8")
    prof = P.load(str(d)); ws = prof.workspace("X_1")
    with open(os.path.join(ws.dir, "checks.yaml"), "w", encoding="utf-8") as f:
        yaml.safe_dump({"system": "a", "ui": [{"id": "U1", "title": "t"}, {"id": "U1", "title": "t2"}]}, f)
    with pytest.raises(SystemExit):
        engine.load(ws)


# ---------------------------------------------------------------- workbook
def _mini_profile(tmp_path, extra=None, areas=None):
    tools = tmp_path / "tools"; tools.mkdir(exist_ok=True)
    tpl = openpyxl.Workbook(); sh = tpl.active
    sh["D3"] = "X_1"; sh["B10"] = "Mục đích"; sh["D10"] = "Kết quả mong muốn"
    tpl.save(tools / "tpl.xlsx")
    d = tmp_path / ".report-kit"; d.mkdir(exist_ok=True)
    cfg = {"name": "demo", "paths": {"tools_dir": str(tools), "work_root": str(tmp_path / "w")},
           "systems": {"a": {"web": "http://a", "api": "http://a/api"}},
           "workbook": {"template": str(tools / "tpl.xlsx"), "out_dir": str(tmp_path / "out"),
                        "coverage": {"areas": {} if areas is None else areas}}}
    cfg.update(extra or {})
    (d / "project.yaml").write_text(yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8")
    return P.load(str(d))


def _function_list(prof, rows, sheet_title="List chức năng"):
    fl = openpyxl.Workbook(); first = fl.active; first.title = "List báo cáo"
    first.append(["Link UC", "UC Name"]); first.append(["OLD", "[# 1E_117] Bản cũ - không dùng"])
    sh = fl.create_sheet(sheet_title)
    sh.append(["KBKT bàn giao"])
    sh.append(["Stt", "Phân hệ", "Mã Jira", "New UC Name", "Use Case Name", "Menu/tab tương ứng", "Menu/tab (cấp thấp nhất)",
               "PIC", "PIC DEV", "Link testcase của DEV"])
    for r in rows:
        sh.append(r)
    fl.save(os.path.join(prof.tools_dir, "list.xlsx"))


FL_CFG = {"inputs": {"function_list": {
    "glob": "list.xlsx", "sheet": "List chức năng", "header_contains": "Mã Jira", "code_columns": ["New UC Name"],
    "columns": {"jira": "Mã Jira", "name": "New UC Name", "system": "Phân hệ", "menu": "Menu/tab (cấp thấp nhất)",
                "pic": "PIC DEV", "testcase": "Link testcase của DEV"},
    "ticket": {"column": "testcase", "regex": "^(AI-\\d+)_"}}}}


def test_function_list_second_sheet(tmp_path):
    from reportkit import inputs
    prof = _mini_profile(tmp_path, FL_CFG)
    _function_list(prof, [
        [117, "IDS", "HSISI-806", "[# 1E_117] Trang nội bộ", "[ # 1E_130] Dashboard", "x", "Trang chủ (dashboard)",
         "Huyền", "Hoang Nghia Tuan", "AI-23_1E117_TrangNoiBo_v4.xlsx"],
        [130, "IDS", "HSISI-819", "[# 1E_130] Vi phạm định kỳ", "[ # 1E_143] Khác", "x", "Thống kê >> Vi phạm CBTT >> Định kỳ",
         "Huyền", "Tran", None]])
    fn = inputs.find_function(prof, "1E_117")
    assert fn["jira"] == "HSISI-806" and fn["pic"] == "Hoang Nghia Tuan" and fn["ticket"] == "AI-23"     # 'PIC DEV' beats 'PIC'
    assert "[List chức năng]" in fn["source"] and fn["name"].startswith("[# 1E_117]")
    fn = inputs.find_function(prof, "1E_130")             # 1E_130 also appears in 'Use Case Name' of row 117: ignored
    assert fn["jira"] == "HSISI-819" and fn["ticket"] is None
    assert inputs.menu_path(fn) == ["Thống kê", "Vi phạm CBTT", "Định kỳ"]
    assert inputs.screen_names(fn) == ["Vi phạm định kỳ"]
    with pytest.raises(SystemExit):
        inputs.find_function(prof, "1E_143")


def test_siblings_share_screen_or_code(tmp_path):
    import json
    from reportkit import inputs
    prof = _mini_profile(tmp_path, FL_CFG)
    _function_list(prof, [[117, "IDS", "J1", "[# 1E_117] Trang nội bộ"], [118, "IDS", "J2", "[# 1E_118] Trang CTĐC"],
                          [119, "IDS", "J3", "[# 1E_119] Báo cáo khác"]])
    root = tmp_path / "w2"
    for code, files, routes in (("1E_117", ["a.ts", "a.html", "z.java"], ["/dashboard"]), ("2E_50", ["a.ts", "a.html", "y.java"], []),
                                ("2E_51", ["a.ts", "q", "r"], [])):
        (root / code).mkdir(parents=True)
        (root / code / "trace.json").write_text(json.dumps({"top": [{"file": f} for f in files], "routes": routes}), encoding="utf-8")
    fn = inputs.find_function(prof, "1E_118")
    sibs = inputs.find_siblings(prof, "1E_118", fn, trace={"top": [{"file": f} for f in ["a.ts", "a.html", "x.java"]], "routes": ["/dashboard"]},
                                workspaces_root=str(root))
    assert [s["code"] for s in sibs] == ["1E_117", "2E_50"]
    assert any("cùng màn hình /dashboard" in w for w in sibs[0]["why"]) and sibs[0]["name"] == "[# 1E_117] Trang nội bộ"
    assert any("2/3" in w for w in sibs[1]["why"])
    assert inputs.find_siblings(prof, "1E_119", inputs.find_function(prof, "1E_119")) == []




def _analysis(ws, text="# a\nR1. Chỉ lấy tin đã duyệt\n"):
    with open(os.path.join(ws.dir, "analysis.md"), "w", encoding="utf-8") as f:
        f.write(text)


def test_workbook_is_one_sheet_and_lists_become_lines(tmp_path):
    from reportkit import workbook as W
    prof = _mini_profile(tmp_path); ws = prof.workspace("X_1"); _analysis(ws)
    wb = W.Workbook(prof, ws, name="n", screen="s", ticket="T", ascii_name="N")
    wb.chapter("Chức năng 1"); wb.cat("Giao diện")
    r = wb.tc("Kiểm tra A", ["Chọn Năm = 2026", "Bấm [Tìm kiếm]"], ["Ý 1", "Ý 2"], basis="Căn cứ: thiết kế chức năng",
              status="F", actual=["Sai.", "Ví dụ: công ty A hiện 0, đúng ra là 45.", "Nguyên nhân: chưa tính lại."], covers="R1")
    wb.finish(run_date="06/10/2026", run_note="n")
    book = openpyxl.load_workbook(wb.save())
    assert len(book.sheetnames) == 1
    sh = book.worksheets[0]
    assert sh.cell(r, 3).value == "1. Chọn Năm = 2026\n2. Bấm [Tìm kiếm]" and sh.cell(r, 4).value == "- Ý 1\n- Ý 2"
    assert sh.cell(r, 8).value == "Sai.\n- Ví dụ: công ty A hiện 0, đúng ra là 45.\n- Nguyên nhân: chưa tính lại."


def test_readability_gate(tmp_path):
    from reportkit import workbook as W
    prof = _mini_profile(tmp_path); ws = prof.workspace("X_1"); _analysis(ws)
    wb = W.Workbook(prof, ws, name="n", screen="s", ticket="T", ascii_name="N")
    wb.chapter("Chức năng 1"); wb.cat("Chức năng")
    wb.tc("Kiểm tra Điểm", "1. Bấm [Tìm kiếm]", "Điểm = 100 trừ điểm bị trừ", status="P",
          actual=["Đạt.", "Ví dụ: công ty BMK kỳ 08/2026 hiện 70, khớp."], basis="Căn cứ: thiết kế chức năng", covers=["R1"])  # clean
    wb.tc("Kiểm tra Phân loại", "1. Bấm [Tìm kiếm]", "A ≥ 90", status="F",
          actual="Sai. 222.032 ô so sánh, vd AMD", bug="BUG-01, BUG-02", basis="Căn cứ: x\nKỹ thuật: D01, EVALUATIONS.TYPE")
    wb.tc("Kiểm tra tổng", "1. Bấm [Tìm kiếm]", "Tổng khớp", status="P", actual="Sai. " + "x" * 200)
    wb.cat("An toàn thông tin")
    wb.tc("Kiểm tra XSS", "1. Nhập <script>alert(1)</script> | ' OR 1=1 --", "Không chạy script", status="P", actual="Đạt. Hiện nguyên văn")
    wb.finish(run_date="06/10/2026", run_note="n")
    got = W.readability(wb.save())
    assert not [x for x in got if x[0] in ("X_1_1", "X_1_4")]                      # clean case; security block skipped
    two = {(c, why.split(":")[0].split(" (")[0]) for cid, c, why in got if cid == "X_1_2"}
    assert ("D", "'≥'") in two and ("H", "'ô so sánh'") in two and ("I", "several bugs in one case") in two
    assert ("J", "run id") in two
    three = [(c, why) for cid, c, why in got if cid == "X_1_3"]
    assert any(c == "H" and why.startswith("must start with 'Đạt.'") for c, why in three)
    assert any(c == "H" and "160" in why for c, why in three)


def test_coverage_gate_blocks_until_every_requirement_has_a_case_or_reason(tmp_path):
    from reportkit import workbook as W
    areas = {"bo_loc": "Từng bộ lọc", "xss": "XSS"}
    prof = _mini_profile(tmp_path, areas=areas); ws = prof.workspace("X_1")
    _analysis(ws, "# Quy tắc\nR1. Chỉ lấy tin đã duyệt\n- R2: Bản đính chính mới nhất thay bản gốc\n"
                  "| R3 | Hạn nộp theo ngày làm việc |\nKhông phải quy tắc: (R4) nằm giữa dòng\n")
    with open(os.path.join(ws.dir, "probe.md"), "w", encoding="utf-8") as f:
        f.write("## T\n- [FOUND] lang_twins: Bản VI và EN song song\n- [none] orphan_company: x\n")
    req = W.requirements(ws.dir, areas)
    assert list(req) == ["R1", "R2", "R3", "trap:lang_twins", "area:bo_loc", "area:xss"]

    def build(version, waive_all):
        wb = W.Workbook(prof, ws, name="n", screen="s", ticket="T", ascii_name="N", version=version)
        wb.chapter("Chức năng 1"); wb.cat("Chức năng")
        wb.tc("Kiểm tra trạng thái", "1. Bấm [Tìm kiếm]", "Chỉ tin đã duyệt", status="P", actual="Đạt. Khớp.",
              covers=["R1", "trap:lang_twins", "area:bo_loc"])
        wb.tc("Kiểm tra hạn nộp", "1. Bấm [Tìm kiếm]", "Hạn đúng", actual="Chưa thực hiện - cần dữ liệu", covers="R3")
        if waive_all:
            wb.waive("R2", "Không có chuỗi đính chính trên môi trường")
            wb.waive("area:xss", "Màn hình không có ô nhập chữ nào")
        wb.finish(run_date="07/10/2026", run_note="n")
        return wb.save()

    with pytest.raises(SystemExit) as e:
        build(1, waive_all=False)
    assert "R2" in str(e.value) and "area:xss" in str(e.value) and "No file written" in str(e.value)
    assert not os.path.exists(tmp_path / "out" / "T_X1_N_v1.xlsx")                # the version is not used up
    out = build(1, waive_all=True)
    assert os.path.exists(out)
    cov = open(os.path.join(ws.dir, "coverage.md"), encoding="utf-8").read()
    assert "6 requirements, 4 by cases, 2 without a case (reason), 0 missing" in cov
    assert "X_1_1 Kiểm tra trạng thái" in cov and "Không có case: Không có chuỗi đính chính" in cov

    wb = W.Workbook(prof, ws, name="n", screen="s", ticket="T", ascii_name="N", version=2)
    wb.tc("Kiểm tra", "1. a", "b", status="P", actual="Đạt. x",
          covers=["R1", "R2", "R3", "trap:lang_twins", "area:bo_loc", "area:xss", "R9"])
    with pytest.raises(SystemExit) as e:                                         # typo / unknown id
        wb.save()
    assert "unknown ids" in str(e.value) and "R9" in str(e.value)
    with pytest.raises(ValueError):
        wb.waive("R2", "")


def test_coverage_needs_numbered_analysis(tmp_path):
    from reportkit import workbook as W
    prof = _mini_profile(tmp_path); ws = prof.workspace("X_1")
    with pytest.raises(SystemExit):
        W.requirements(ws.dir)                                                   # no analysis.md
    _analysis(ws, "# a\nChỉ lấy tin đã duyệt, không đánh số\n")
    with pytest.raises(SystemExit):
        W.requirements(ws.dir)


def test_build_check_writes_scratch_file(tmp_path, monkeypatch):
    from reportkit import workbook as W
    prof = _mini_profile(tmp_path); ws = prof.workspace("X_1"); _analysis(ws)
    scratch = os.path.join(ws.dir, "_check.xlsx")
    monkeypatch.setenv("RK_WB_OUT", scratch)
    for _ in range(2):                                                           # re-runnable: the scratch file is replaced
        wb = W.Workbook(prof, ws, name="n", screen="s", ticket="T", ascii_name="N", version=1)
        wb.tc("Kiểm tra", "1. a", "b", status="P", actual="Đạt. x", covers="R1")
        wb.finish(run_date="07/10/2026", run_note="n")
        assert wb.save() == scratch
    assert not os.path.exists(tmp_path / "out")


def test_summary_pick_blocks_by_id_and_status():
    from reportkit.checks import engine
    text = "# Run r1 - 1D_110 (ids)\nCounts: DIFF=1, MATCH=1, OBS=1\n\n## D01 [MATCH] a\ncells compared: 4, differing: 0\n\n" \
           "## D02 [DIFF] b\n- x: DIFF\n\n## U01 [OBS] c  (carried from r0)\n- title: \"T\"\n"
    out, n = engine.pick_blocks(text, ["U01"], ["diff"])
    assert n == 2 and "## D02" in out and "## U01" in out and "## D01" not in out and out.startswith("# Run r1")
    out, n = engine.pick_blocks(text, [""], [""])
    assert n == 3
